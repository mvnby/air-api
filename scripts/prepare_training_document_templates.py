"""Create deterministic synthetic DOCX sources for the sanctioned native bundle importer.

This writes local template source files only. It has no DB or network access.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


BUNDLE_ID = "kitlane-training-v1"
TEMPLATES = (
    ("offer", "Коммерческое предложение", None),
    ("invoice", "Счёт на оплату", None),
    ("contract", "Договор", None),
)


def template_bytes(doc_type: str, title: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name in ("Title", "Heading 1"):
        style = document.styles[name]
        style.font.name, style.font.color.rgb = "Arial", RGBColor(0, 0, 0)
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    document.styles["Title"].font.size = Pt(21)
    marker = document.add_paragraph("УЧЕБНЫЙ ОБРАЗЕЦ")
    marker.runs[0].bold = True
    document.add_paragraph(title, "Title")
    document.add_paragraph(
        "Для демонстрации работы CRM Kitlane. Все участники вымышлены. "
        "Образец не предназначен для подписания, оплаты или отправки клиентам."
    )
    document.add_paragraph("№ {{document.official_full_number}} от {{document.issued_on}}")
    document.add_paragraph("Номер CRM: {{document.internal_reference}}")
    document.add_paragraph("Продавец: {{seller.legal_name}}")
    document.add_paragraph("Клиент: {{customer.full_name}}")
    document.add_paragraph("Оборудование", "Heading 1")
    table = document.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = (75, 30, 35, 30)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Mm(width)
    for cell, text in zip(table.rows[0].cells, ("Наименование", "Кол-во", "Цена BYN", "Сумма BYN"), strict=True):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for cell, text in zip(table.rows[1].cells,
                          ("{{lines}}{{line.title}}", "{{line.quantity}}", "{{line.unit_price}}", "{{line.amount}}"), strict=True):
        cell.text = text
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    document.add_paragraph("Итого: {{totals.amount}} {{totals.currency}}")
    document.add_paragraph("Учебные условия", "Heading 1")
    descriptions = {
        "offer": "Сравните варианты оборудования в заказе и выберите подходящее предложение. Указанные цены служат только для обучения.",
        "invoice": "Учебный счёт показывает выбранные строки заказа. Банковские реквизиты отсутствуют; оплачивать этот образец не требуется.",
        "contract": "Учебный макет показывает стороны, предмет и сумму сделки. Создайте черновик в документах заказа, проверьте PDF и выпустите учебный образец с номером. Юридическая полнота этого образца не заявляется.",
    }
    document.add_paragraph(descriptions[doc_type])
    document.add_paragraph("Подписи и печати отсутствуют. УЧЕБНЫЙ ОБРАЗЕЦ.")
    document.core_properties.author = "Kitlane training"
    document.core_properties.title = f"УЧЕБНЫЙ ОБРАЗЕЦ {title}"
    document.core_properties.created = document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    raw = BytesIO()
    document.save(raw)
    # python-docx ZIP timestamps otherwise make hashes vary on every run.
    normalized = BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as source, ZipFile(normalized, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return normalized.getvalue()


def prepare(output_dir: Path) -> dict:
    output_dir.mkdir(mode=0o700, parents=True, exist_ok=True)
    manifest = {"bundle_id": BUNDLE_ID, "description": "Три полностью вымышленных учебных образца", "templates": []}
    for index, (doc_type, title, scenario) in enumerate(TEMPLATES):
        content = template_bytes(doc_type, title)
        filename = f"training-{doc_type}.docx"
        destination = output_dir / filename
        if destination.exists() and destination.read_bytes() != content:
            raise ValueError(f"Existing source differs; use a new private output directory: {filename}")
        destination.write_bytes(content)
        destination.chmod(0o600)
        manifest["templates"].append({
            "key": f"training-{doc_type}", "name": f"[ОБУЧЕНИЕ] {title}",
            "aliases": [], "doc_type": doc_type, "contract_scenario": scenario,
            "source_filename": filename, "sha256": sha256(content).hexdigest(),
            "description": "УЧЕБНЫЙ ОБРАЗЕЦ. Вымышленные данные; без подписей и банковских реквизитов.",
            "sort_order": (index + 1) * 10,
        })
    manifest_path = output_dir / "bundle.json"
    encoded = json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    if manifest_path.exists() and manifest_path.read_text() != encoded:
        raise ValueError("Existing manifest differs; use a new private output directory")
    manifest_path.write_text(encoded, encoding="utf-8")
    manifest_path.chmod(0o600)
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(prepare(args.output_dir), ensure_ascii=False, indent=2))
