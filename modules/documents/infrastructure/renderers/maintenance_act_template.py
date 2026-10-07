"""Neutral factual template; version 1 must stay reproducible for saved drafts."""
from io import BytesIO
from datetime import datetime

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


TYPE = "maintenance_defect_act"
NAME = "Дефектный акт по замечаниям ТО · v1"
FIELDS = frozenset({"document.official_full_number", "document.issued_on", "seller.legal_name",
                    "customer.full_name", "order.object_title", "order.object_address", "maintenance.source_order"})
ROW_FIELDS = frozenset({"observation.equipment", "observation.provenance", "observation.facts",
                        "observation.recommendation", "observation.sources", "observation.photos"})


def template_bytes():
    document = Document()
    document.core_properties.author = ""
    document.core_properties.created = datetime(2026, 10, 8)
    document.core_properties.modified = document.core_properties.created
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Normal'].font.size = Pt(10)
    document.styles['Normal'].paragraph_format.space_after = Pt(5)
    document.add_heading("Дефектный акт № {{ document.official_full_number }}", 1)
    document.add_paragraph("Дата: {{ document.issued_on }}")
    document.add_paragraph("Исполнитель: {{ seller.legal_name }}")
    document.add_paragraph("Заказчик: {{ customer.full_name }}")
    document.add_paragraph("Объект: {{ order.object_title }}\nАдрес: {{ order.object_address }}")
    document.add_paragraph("Исходное ТО: {{ maintenance.source_order }}")
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(17)
    cell = table.cell(0, 0)
    cell.vertical_alignment = 1
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            element.set(qn(f"w:{key}"), value)
        borders.append(element)
    properties.append(borders)
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), "140"); element.set(qn("w:type"), "dxa"); margins.append(element)
    properties.append(margins)
    cell.paragraphs[0].text = "{{ observations }}{{ observation.equipment }}"
    cell.paragraphs[0].runs[0].bold = True
    for label, field in (("", "provenance"), ("Подтверждённые факты", "facts"),
                         ("Рекомендация (предстоящие действия)", "recommendation"),
                         ("Источник", "sources"), ("Приватные фото", "photos")):
        paragraph = cell.add_paragraph()
        if label:
            paragraph.add_run(label + ": ").bold = True
        paragraph.add_run("{{ observation." + field + " }}")
    output = BytesIO(); document.save(output)
    return output.getvalue()
