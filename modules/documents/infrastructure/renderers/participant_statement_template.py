"""Editable generic statement; v1 remains reproducible for existing drafts."""
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt, RGBColor

TYPE = "participant_statement"
NAME = "Заявление участника · редактируемая форма v1"
FIELDS = frozenset({
    "document.official_full_number", "document.issued_on", "seller.legal_name", "seller.unp",
    "seller.signer_position", "seller.signer_name", "seller.acting_basis",
    "participant.procedure_reference", "participant.lot", "participant.buyer_name",
    "participant.declaration_text",
})


def template_bytes() -> bytes:
    document = Document()
    document.core_properties.author = ""
    document.core_properties.created = datetime(2026, 10, 9)
    document.core_properties.modified = document.core_properties.created
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ("Normal", "Title"):
        style = document.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for border in document.styles["Title"].element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.styles["Title"].font.size = Pt(16)
    document.add_paragraph("Заявление участника о соответствии", style="Title")
    document.add_paragraph("№ {{ document.official_full_number }} от {{ document.issued_on }}")
    document.add_paragraph("Заказчик закупки: {{ participant.buyer_name }}")
    document.add_paragraph("Процедура закупки: {{ participant.procedure_reference }}")
    document.add_paragraph("Лот: {{ participant.lot }}")
    document.add_paragraph("Участник: {{ seller.legal_name }}\nУНП: {{ seller.unp }}")
    document.add_paragraph("{{ participant.declaration_text }}")
    document.add_paragraph("Подписант: {{ seller.signer_position }} {{ seller.signer_name }}")
    document.add_paragraph("{{ seller.acting_basis }}")
    document.add_paragraph("Подпись: ____________________")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
