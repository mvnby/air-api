"""Selected-template customer requirements, based only on frozen factual values."""

from __future__ import annotations

import asyncio
from io import BytesIO
from collections.abc import Mapping
import re

from docx import Document
from docx.shared import Pt, RGBColor

from modules.documents.domain.party import INDIVIDUAL, INDIVIDUAL_ENTREPRENEUR, SELF
from modules.documents.domain.placeholder_catalog import SCALAR_PLACEHOLDERS
from modules.documents.infrastructure.renderers import NativeDocxRenderer
from .context_builder import DocumentContextBuilder
from .errors import ManagedDocumentConflictError, ManagedDocumentNotFoundError
from .template_selection import (
    resolve_template_use_case,
    select_active_native_template,
    load_document_template_version,
)


READINESS_DOCUMENT_TYPES = frozenset({"contract", "invoice"})
CRITICAL_CUSTOMER_FIELDS = frozenset(
    {
        "customer.display_name",
        "customer.full_name",
        "customer.unp",
        "customer.legal_address",
        "customer.bank_name",
        "customer.iban",
        "customer.bic",
        "customer.signer_position",
        "customer.signer_name",
        "customer.acting_basis",
    }
)
OPTIONAL_CUSTOMER_FIELDS = frozenset(
    {"customer.phone", "customer.email", "customer.city"}
)
LABELS = {item.name: item.label for item in SCALAR_PLACEHOLDERS}


def _present(value) -> bool:
    text = str(value or "").strip()
    return bool(text) and not re.match(r"^_{3,}(?:\s|$)", text)


def snapshot_customer_readiness(
    snapshot: dict, *, source: bytes, document_type: str
) -> dict:
    if document_type not in READINESS_DOCUMENT_TYPES:
        return {"checked": False, "missing_fields": [], "can_issue": True}
    if not isinstance(snapshot, Mapping) or not isinstance(
        snapshot.get("values"), Mapping
    ):
        raise ValueError(
            "Сохранённый снимок документа повреждён: отсутствует каталог фактических значений"
        )
    values = snapshot["values"]
    for field in CRITICAL_CUSTOMER_FIELDS | OPTIONAL_CUSTOMER_FIELDS:
        if values.get(field) is not None and not isinstance(values[field], str):
            raise ValueError(
                "Сохранённый снимок документа повреждён: поле клиента должно быть строкой: "
                + LABELS[field]
            )
    if not isinstance(snapshot.get("conditions", {}), Mapping):
        raise ValueError(
            "Сохранённый снимок документа повреждён: каталог условий должен быть объектом"
        )
    entity = values.get("customer.entity_type")
    mode = values.get("customer.signing_mode")
    effective = NativeDocxRenderer().discover_active_placeholders(
        source, snapshot.get("conditions") or {}
    )
    applicable = set(CRITICAL_CUSTOMER_FIELDS | OPTIONAL_CUSTOMER_FIELDS)
    if entity == INDIVIDUAL:
        applicable -= {
            "customer.unp",
            "customer.bank_name",
            "customer.iban",
            "customer.bic",
        }
    if entity in {INDIVIDUAL, INDIVIDUAL_ENTREPRENEUR}:
        applicable.discard("customer.signer_position")
    if mode == SELF:
        applicable -= {
            "customer.signer_position",
            "customer.signer_name",
            "customer.acting_basis",
        }
    missing = [
        {
            "field": field,
            "label": (
                "Адрес клиента"
                if entity == INDIVIDUAL and field == "customer.legal_address"
                else "ФИО клиента"
                if entity == INDIVIDUAL
                and field in {"customer.full_name", "customer.display_name"}
                else LABELS[field]
            ),
            "critical": field in CRITICAL_CUSTOMER_FIELDS,
        }
        for field in sorted(effective & applicable)
        if not _present(values.get(field))
    ]
    return {
        "checked": True,
        "missing_fields": missing,
        "can_issue": not any(item["critical"] for item in missing),
    }


async def source_customer_readiness(
    snapshot, *, template, version, template_storage, document_type
):
    if document_type not in READINESS_DOCUMENT_TYPES:
        return {"checked": False, "missing_fields": [], "can_issue": True}
    if template_storage is None:
        raise ManagedDocumentConflictError(
            "Для проверки реквизитов требуется исходный DOCX шаблона"
        )
    source = await template_storage.read_persisted(
        tenant_id=template.tenant_id,
        template_id=template.id,
        version=version.version,
        storage_key=version.source_storage_key,
        filename=str(version.source_filename or "template.docx"),
        checksum_sha256=version.checksum_sha256,
    )
    return await asyncio.to_thread(
        snapshot_customer_readiness,
        snapshot,
        source=source,
        document_type=document_type,
    )


async def check_selection_readiness(
    session, *, tenant_scope, selection, template_id, template_storage
):
    from .lifecycle_service import ManagedDocumentService

    order = await ManagedDocumentService._get_mutable_scoped_order(
        session,
        tenant_scope=tenant_scope,
        order_id=selection.order_id,
        require_mutable=True,
    )
    if order is None:
        raise ManagedDocumentNotFoundError("Заказ не найден")
    scenario, role = resolve_template_use_case(selection)
    template, version = await select_active_native_template(
        session,
        tenant_scope=tenant_scope,
        legal_entity_id=selection.legal_entity_id,
        document_type=selection.document_type,
        template_id=template_id,
        contract_scenario=scenario,
        business_role=role,
    )
    snapshot = await DocumentContextBuilder.build(
        session,
        tenant_scope=tenant_scope,
        selection=selection,
        template=template,
        template_version=version,
        template_storage=template_storage,
    )
    readiness = await source_customer_readiness(
        snapshot,
        template=template,
        version=version,
        template_storage=template_storage,
        document_type=selection.document_type,
    )
    return {**readiness, "template_id": template.id, "template_version_id": version.id}


async def check_saved_document_readiness(
    session, *, tenant_scope, document, template_storage
):
    if document.doc_type not in READINESS_DOCUMENT_TYPES or document.status != "draft":
        return {"checked": False, "missing_fields": [], "can_issue": True}
    template, version = await load_document_template_version(
        session, tenant_scope=tenant_scope, document=document
    )
    return await source_customer_readiness(
        document.render_snapshot or {},
        template=template,
        version=version,
        template_storage=template_storage,
        document_type=document.doc_type,
    )


def assert_customer_ready(readiness: dict) -> None:
    missing = [
        item["label"]
        for item in readiness.get("missing_fields", [])
        if item.get("critical")
    ]
    if missing:
        raise ManagedDocumentConflictError(
            "Не заполнены поля клиента: "
            + ", ".join(missing)
            + ". Заполните карточку и создайте новый черновик; сохранённый снимок не изменяется."
        )


INCOMPLETE_DRAFT_MARKER = "ЧЕРНОВИК — не заполнены обязательные поля клиента"


def mark_incomplete_docx(content: bytes, readiness: dict) -> bytes:
    """Mark the output copy, including every active section/header variant."""
    if readiness.get("can_issue", True):
        return content
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError("Черновик DOCX повреждён или не читается") from exc

    def add_marker(container):
        if any(p.text == INCOMPLETE_DRAFT_MARKER for p in container.paragraphs):
            return
        paragraph = (
            container.paragraphs[0].insert_paragraph_before()
            if container.paragraphs
            else container.add_paragraph()
        )
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(INCOMPLETE_DRAFT_MARKER)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(160, 30, 30)

    add_marker(document)
    seen = set()
    for section in document.sections:
        headers = [section.header]
        if section.different_first_page_header_footer:
            headers.append(section.first_page_header)
        if document.settings.odd_and_even_pages_header_footer:
            headers.append(section.even_page_header)
        for header in headers:
            element = header._element
            if element in seen:
                continue
            seen.add(element)
            add_marker(header)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
