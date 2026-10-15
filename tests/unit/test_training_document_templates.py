from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pytest

from modules.documents.infrastructure.renderers import DocumentTemplateVersion, NativeDocxRenderer, RenderContext
from scripts.manage_native_document_template_bundle import load_bundle, prepare_sources
from scripts.prepare_training_document_templates import prepare


def test_training_sources_are_deterministic_safe_and_render_two_product_rows(tmp_path):
    first, second = tmp_path / "first", tmp_path / "second"
    assert prepare(first) == prepare(second)
    committed = Path(__file__).resolve().parents[2] / "config/native_document_template_bundles/kitlane-training-v1.json"
    assert json.loads(committed.read_text()) == json.loads((first / "bundle.json").read_text())
    bundle = load_bundle(first / "bundle.json")
    sources = prepare_sources(bundle, first)
    assert {spec.doc_type for spec in bundle.templates} == {"offer", "invoice", "contract"}
    for spec in bundle.templates:
        content = sources[spec.key].content
        assert content == (second / spec.source_filename).read_bytes()
        with ZipFile(BytesIO(content)) as archive:
            assert not any(name.startswith("word/media/") for name in archive.namelist())
            assert b"pBdr" not in archive.read("word/styles.xml")
        contract = sources[spec.key].contract
        template = DocumentTemplateVersion(template_key=spec.key, version=1, source=content,
            field_catalog=contract.field_catalog, condition_catalog=contract.condition_catalog,
            table_blocks=contract.table_blocks)
        assert NativeDocxRenderer().validate(template).is_valid
        rendered = NativeDocxRenderer().render(template, RenderContext(values={
            "document.issued_on": "15.10.2026", "document.internal_reference": "УЧ-001",
            "document.official_full_number": "УЧ-2026-001",
            "seller.legal_name": "[ОБУЧЕНИЕ] Учебная компания Kitlane",
            "customer.full_name": "[ОБУЧЕНИЕ] ООО Пример", "totals.amount": "5310,00",
            "totals.currency": "BYN",
        }, table_rows={"lines": (
            {"line.title": "Учебная модель 1", "line.quantity": "1", "line.unit_price": "2490,00", "line.amount": "2490,00"},
            {"line.title": "Учебная модель 2", "line.quantity": "1", "line.unit_price": "2820,00", "line.amount": "2820,00"},
        )}))
        document = Document(BytesIO(rendered.content))
        assert len(document.tables[0].rows) == 3
        text = "\n".join([p.text for p in document.paragraphs] + [c.text for r in document.tables[0].rows for c in r.cells])
        assert "УЧЕБНЫЙ ОБРАЗЕЦ" in text and "{{" not in text
        assert "№ УЧ-2026-001 от 15.10.2026" in text
        assert "Учебная модель 1" in text and "Учебная модель 2" in text


def test_modified_source_is_not_silently_overwritten(tmp_path):
    prepare(tmp_path)
    (tmp_path / "training-offer.docx").write_bytes(b"changed by user")
    with pytest.raises(ValueError, match="Existing source differs"):
        prepare(tmp_path)
