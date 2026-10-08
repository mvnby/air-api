from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from modules.documents.application.artifact_helpers import _condition_values, build_render_inputs
from modules.documents.infrastructure.renderers import NativeDocxRenderer



def test_condition_values_require_complete_boolean_snapshot() -> None:
    catalog = frozenset(
        {
            "seller.is_individual_entrepreneur",
            "customer.is_organization",
        }
    )

    with pytest.raises(ValueError, match="отсутствуют условные флаги"):
        _condition_values(
            catalog,
            {"seller.is_individual_entrepreneur": True},
        )
    with pytest.raises(TypeError, match="должен быть boolean"):
        _condition_values(
            catalog,
            {
                "seller.is_individual_entrepreneur": "false",
                "customer.is_organization": True,
            },
        )


def test_condition_values_keep_old_conditionless_templates_compatible() -> None:
    assert _condition_values(frozenset(), None) == {}


def _signing_render_inputs(document, snapshot):
    buffer = BytesIO()
    document.save(buffer)
    return build_render_inputs(
        template=SimpleNamespace(id=1),
        version=SimpleNamespace(
            version=1, source_filename="contract.docx",
            placeholder_schema={"fields": ["customer.signer_position", "customer.acting_basis"]},
        ),
        source=buffer.getvalue(), snapshot=snapshot,
    )


@pytest.mark.parametrize(
    "entity,mode,expected_position,expected_basis",
    [("organization", "statutory_body", True, True),
     ("organization", "power_of_attorney", True, True),
     ("individual_entrepreneur", "self", False, False),
     ("individual", "self", False, False),
     ("individual_entrepreneur", "power_of_attorney", False, True),
     ("individual", "power_of_attorney", False, True)],
)
def test_native_customer_signing_lines_are_presentation_only(
    entity, mode, expected_position, expected_basis,
):
    from modules.documents.domain.customer_signing import CUSTOMER_POSITION_LINE, CUSTOMER_BASIS_LINE

    document = Document()
    paragraph = document.add_paragraph("В лице ")
    paragraph.add_run("{{ customer.signer_")
    paragraph.add_run("position }}").bold = True
    document.add_paragraph("на основании {{ customer.acting_basis }}")
    document.sections[0].header.paragraphs[0].text = "{{ customer.signer_position }}"
    document.sections[0].footer.paragraphs[0].text = "{{ customer.acting_basis }}"
    document.add_table(rows=1, cols=1).cell(0, 0).text = "{{ customer.acting_basis }}"
    values = {
        "customer.entity_type": entity, "customer.signing_mode": mode,
        "customer.signer_position": "  ", "customer.acting_basis": "",
    }
    snapshot = {"values": values}
    template, context = _signing_render_inputs(document, snapshot)
    output = Document(BytesIO(NativeDocxRenderer().render(template, context).content))
    position = CUSTOMER_POSITION_LINE if expected_position else ""
    basis = CUSTOMER_BASIS_LINE if expected_basis else ""
    assert output.paragraphs[0].text == "В лице " + position
    assert output.paragraphs[1].text == "на основании " + basis
    assert output.sections[0].header.paragraphs[0].text == position
    assert output.sections[0].footer.paragraphs[0].text == basis
    assert output.tables[0].cell(0, 0).text == basis
    assert "customer.entity_type" not in context.values
    assert "customer.signing_mode" not in context.values
    assert snapshot["values"]["customer.signer_position"] == "  "
    assert snapshot["values"]["customer.acting_basis"] == ""


def test_native_known_signing_facts_keep_text_and_formatting():
    values = {"customer.entity_type": "organization", "customer.signing_mode": "power_of_attorney",
              "customer.signer_position": "представителя", "customer.acting_basis": "Доверенности № 7"}
    document = Document()
    document.add_paragraph().add_run("{{ customer.signer_position }}").bold = True
    document.add_paragraph("{{ customer.acting_basis }}")
    template, context = _signing_render_inputs(document, {"values": values})
    output = Document(BytesIO(NativeDocxRenderer().render(template, context).content))
    assert [p.text for p in output.paragraphs] == ["представителя", "Доверенности № 7"]
    assert output.paragraphs[0].runs[0].bold is True
