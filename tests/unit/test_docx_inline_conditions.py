"""Inline conditions must preserve Word structure and remain paragraph-local."""
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from modules.documents.infrastructure.renderers import (
    DocumentTemplateVersion,
    NativeDocxRenderer,
    RenderContext,
    TemplateValidationError,
)


SELLER = "seller.individual_entrepreneur_self"
CUSTOMER = "customer.organization_statutory_body"


def render(document, conditions, values=None):
    source = BytesIO()
    document.save(source)
    template = DocumentTemplateVersion(
        template_key="contract", version=1, source=source.getvalue(),
        field_catalog=frozenset({"seller.city"}),
        condition_catalog=frozenset({SELLER, CUSTOMER}),
    )
    result = NativeDocxRenderer().render(
        template, RenderContext(values=values or {}, conditions=conditions)
    )
    return Document(BytesIO(result.content))


@pytest.mark.parametrize("seller", [False, True])
@pytest.mark.parametrize("customer", [False, True])
def test_inline_preamble_preserves_runs_quotes_and_single_paragraph(seller, customer):
    document = Document()
    p = document.add_paragraph()
    p.add_run("{{#if seller.individual_")
    p.add_run('entrepreneur_self}}ИП, “')
    p.add_run("Исполнитель").bold = True
    p.add_run(f'”, {{{{ seller.city }}}}{{{{/if {SELLER}}}}}')
    p.add_run(f'{{{{#if {CUSTOMER}}}}} и ООО, «')
    p.add_run("Заказчик").bold = True
    p.add_run(f'»{{{{/if {CUSTOMER}}}}}. Договор.')
    result = render(document, {SELLER: seller, CUSTOMER: customer}, {"seller.city": "Минск"})
    assert len(result.paragraphs) == 1
    assert result.paragraphs[0].text == (
        ('ИП, “Исполнитель”, Минск' if seller else '')
        + (' и ООО, «Заказчик»' if customer else '') + '. Договор.'
    )
    assert ''.join(r.text for r in result.paragraphs[0].runs if r.bold) == (
        ('Исполнитель' if seller else '') + ('Заказчик' if customer else '')
    )


@pytest.mark.parametrize("outer,inner", [(False, False), (False, True), (True, False), (True, True)])
def test_nested_inline_conditions_inside_block(outer, inner):
    document = Document()
    document.add_paragraph(f"{{{{#if {SELLER}}}}}")
    document.add_paragraph(f"A{{{{#if {SELLER}}}}}B{{{{#if {CUSTOMER}}}}}C{{{{/if {CUSTOMER}}}}}D{{{{/if {SELLER}}}}}E")
    document.add_paragraph(f"{{{{/if {SELLER}}}}}")
    document.add_paragraph("End")
    result = render(document, {SELLER: outer, CUSTOMER: inner})
    assert [p.text for p in result.paragraphs] == (
        ["AB" + ("C" if inner else "") + "DE", "End"] if outer else ["End"]
    )


@pytest.mark.parametrize("enabled", [False, True])
def test_inline_conditions_in_table_nested_table_and_header(enabled):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 1).text = "Other cell"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    areas = [table.cell(0, 0).paragraphs[0], nested.cell(0, 0).paragraphs[0]]
    for name in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
        areas.append(getattr(document.sections[0], name).paragraphs[0])
    for p in areas:
        p.text = f"Before {{{{#if {SELLER}}}}}inside{{{{/if {SELLER}}}}} after"
    result = render(document, {SELLER: enabled})
    expected = "Before " + ("inside" if enabled else "") + " after"
    assert result.tables[0].cell(0, 0).paragraphs[0].text == expected
    assert result.tables[0].cell(0, 0).tables[0].cell(0, 0).text == expected
    assert result.tables[0].cell(0, 1).text == "Other cell"
    for name in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
        assert getattr(result.sections[0], name).paragraphs[0].text == expected


@pytest.mark.parametrize("enabled", [False, True])
def test_inline_removes_hidden_breaks_tabs_and_preserves_hyperlink(enabled):
    document = Document()
    p = document.add_paragraph(f"Before{{{{#if {SELLER}}}}}")
    p.add_run("\n\t")
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "target")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "link"
    run.append(text)
    link.append(run)
    p._p.append(link)
    p.add_run(f"{{{{/if {SELLER}}}}} after")
    result = render(document, {SELLER: enabled})
    p = result.paragraphs[0]
    assert len(list(p._p.iter(qn("w:br")))) == int(enabled)
    assert len(list(p._p.iter(qn("w:tab")))) == int(enabled)
    links = list(p._p.iter(qn("w:hyperlink")))
    assert links[0].get(qn("w:anchor")) == "target"
    assert ''.join(n.text or '' for n in links[0].iter(qn("w:t"))) == ("link" if enabled else "")


@pytest.mark.parametrize("texts,code", [
    ([f"A{{{{#if {SELLER}}}}}B", f"C{{{{/if {SELLER}}}}}D"], "condition_marker_placement"),
    ([f"A{{{{#if {SELLER}}}}}B{{{{/if {CUSTOMER}}}}}C"], "mismatched_condition_marker"),
    (["A{{#if seller.unknown}}B{{/if seller.unknown}}C"], "unknown_condition"),
    ([f"A{{{{/if {SELLER}}}}}B"], "unbalanced_condition_marker"),
    (["A{{#if seller == ip}}B{{/if seller == ip}}C"], "malformed_condition_marker"),
])
def test_invalid_inline_conditions_fail_closed(texts, code):
    document = Document()
    for text in texts:
        document.add_paragraph(text)
    with pytest.raises(TemplateValidationError) as error:
        render(document, {SELLER: True, CUSTOMER: True})
    assert code in {issue.code for issue in error.value.result.issues}


@pytest.mark.parametrize("enabled", [False, True])
@pytest.mark.parametrize("with_text", [False, True])
def test_inline_controls_at_closing_boundary(enabled, with_text):
    document = Document()
    p = document.add_paragraph("Before" + "{{#if " + SELLER + "}}")
    if with_text:
        p.add_run("inside")
    p.add_run("\n\t")
    p.add_run("{{/if " + SELLER + "}}" + "after")
    result = render(document, {SELLER: enabled})
    assert result.paragraphs[0].text == (
        "Before" + (("inside" if with_text else "") + "\n\t" if enabled else "") + "after"
    )
