from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START

from modules.documents.application.artifact_helpers import build_render_inputs
from modules.documents.application.customer_readiness import (
    INCOMPLETE_DRAFT_MARKER,
    check_saved_document_readiness,
    mark_incomplete_docx,
    snapshot_customer_readiness,
)
from modules.documents.domain.party import party_conditions
from modules.documents.infrastructure.renderers import NativeDocxRenderer


def source_document(*paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    return document


def source_bytes(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def snapshot(entity="organization", mode="statutory_body", **values):
    return {
        "values": {
            "document.type": "contract",
            "customer.entity_type": entity,
            "customer.signing_mode": mode,
            **values,
        },
        "conditions": party_conditions(
            "customer", entity_type=entity, signing_mode=mode
        ),
        "meta": {},
    }


def fields(result):
    return {item["field"] for item in result["missing_fields"]}


@pytest.mark.parametrize(
    "entity,mode,expected",
    [
        (
            "organization",
            "statutory_body",
            {
                "customer.signer_name",
                "customer.signer_position",
                "customer.acting_basis",
            },
        ),
        (
            "organization",
            "power_of_attorney",
            {"customer.signer_name", "customer.acting_basis"},
        ),
        (
            "individual_entrepreneur",
            "power_of_attorney",
            {"customer.signer_name", "customer.acting_basis"},
        ),
        (
            "individual",
            "power_of_attorney",
            {"customer.signer_name", "customer.acting_basis"},
        ),
        ("individual_entrepreneur", "self", set()),
        ("individual", "self", set()),
    ],
)
def test_only_effective_frozen_signing_branch_is_required(entity, mode, expected):
    source = source_bytes(
        source_document(
            "{{#if customer.organization_statutory_body}}{{ customer.signer_position }} "
            "{{ customer.signer_name }} {{ customer.acting_basis }}{{/if customer.organization_statutory_body}}"
            "{{#if customer.signs_by_power_of_attorney}}{{ customer.signer_name }} {{ customer.acting_basis }}{{/if customer.signs_by_power_of_attorney}}"
            "{{#if customer.signs_self}}Лично{{/if customer.signs_self}}"
        )
    )
    frozen = snapshot(entity, mode)
    original = deepcopy(frozen)
    result = snapshot_customer_readiness(
        frozen, source=source, document_type="contract"
    )
    assert fields(result) == expected
    assert result["can_issue"] is not bool(expected)
    assert frozen == original


@pytest.mark.parametrize("empty", [None, "", "  ", "________________ (поле)"])
def test_missing_facts_in_body_table_header_footer_and_optional_contacts(empty):
    document = source_document("{{ customer.full_name }}")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "{{ customer.unp }}"
    document.sections[0].header.paragraphs[0].text = "{{ customer.legal_address }}"
    document.sections[0].footer.paragraphs[0].text = "{{ customer.email }}"
    frozen = snapshot(**{"customer.full_name": "Известное имя", "customer.unp": empty})
    result = snapshot_customer_readiness(
        frozen, source=source_bytes(document), document_type="invoice"
    )
    assert fields(result) == {
        "customer.unp",
        "customer.legal_address",
        "customer.email",
    }
    assert not result["can_issue"]
    assert (
        next(
            item
            for item in result["missing_fields"]
            if item["field"] == "customer.email"
        )["critical"]
        is False
    )


def test_empty_card_does_not_block_template_without_required_customer_fields():
    frozen = snapshot()
    source = source_bytes(source_document("{{ document.type }} {{ customer.email }}"))
    result = snapshot_customer_readiness(
        frozen, source=source, document_type="contract"
    )
    assert result["can_issue"] and fields(result) == {"customer.email"}
    assert mark_incomplete_docx(source, result) == source
    assert (
        snapshot_customer_readiness(frozen, source=source, document_type="act")[
            "checked"
        ]
        is False
    )


def test_historical_literal_basis_is_a_present_fact_without_provenance_guessing():
    result = snapshot_customer_readiness(
        snapshot(**{"customer.acting_basis": "Устава"}),
        source=source_bytes(source_document("{{ customer.acting_basis }}")),
        document_type="contract",
    )
    assert result["can_issue"]


def test_actual_render_ignores_cache_keeps_facts_and_marks_every_header_variant():
    document = source_document(
        "Клиент: {{ customer.full_name }}", "Адрес: {{ customer.legal_address }}"
    )
    section = document.sections[0]
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True
    section.header.paragraphs[0].text = "Обычный колонтитул"
    section.first_page_header.paragraphs[0].text = "Первый колонтитул"
    section.even_page_header.paragraphs[0].text = "Чётный колонтитул"
    document.add_section(WD_SECTION_START.NEW_PAGE).header.is_linked_to_previous = False
    document.sections[1].header.paragraphs[0].text = "Следующий раздел"
    frozen = snapshot(
        **{"customer.full_name": "ООО Клиент", "customer.legal_address": ""}
    )
    frozen["meta"]["customer_readiness"] = {"can_issue": True, "missing_fields": []}
    original = deepcopy(frozen)
    version = SimpleNamespace(
        version=1,
        source_filename="contract.docx",
        placeholder_schema={"fields": ["customer.full_name", "customer.legal_address"]},
    )
    template, context = build_render_inputs(
        template=SimpleNamespace(id=1, doc_type="contract"),
        version=version,
        source=source_bytes(document),
        snapshot=frozen,
    )
    output = NativeDocxRenderer().render(template, context).content
    rendered = Document(BytesIO(output))
    assert frozen == original
    assert "ООО Клиент" in " ".join(p.text for p in rendered.paragraphs)
    assert "____" in context.values["customer.legal_address"]
    assert rendered.paragraphs[0].text == INCOMPLETE_DRAFT_MARKER
    for sec in rendered.sections:
        for header in (sec.header, sec.first_page_header, sec.even_page_header):
            assert (
                sum(p.text == INCOMPLETE_DRAFT_MARKER for p in header.paragraphs) == 1
            )
    marked_again = Document(BytesIO(mark_incomplete_docx(output, {"can_issue": False})))
    assert sum(p.text == INCOMPLETE_DRAFT_MARKER for p in marked_again.paragraphs) == 1
    assert document.sections[0].header.paragraphs[0].text == "Обычный колонтитул"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "metadata",
    [
        None,
        {},
        {"can_issue": True},
        "malformed",
        {"can_issue": True, "missing_fields": []},
    ],
)
async def test_issue_readiness_recomputes_old_stale_malformed_cache_from_exact_saved_source(
    monkeypatch, metadata
):
    import modules.documents.application.customer_readiness as readiness

    frozen = snapshot()
    frozen["meta"]["customer_readiness"] = metadata
    doc = SimpleNamespace(status="draft", doc_type="contract", render_snapshot=frozen)
    template = SimpleNamespace(id=7, tenant_id=3)
    version = SimpleNamespace(
        version=4,
        source_storage_key="frozen-key",
        source_filename="frozen.docx",
        checksum_sha256="sha",
    )

    async def load(*args, **kwargs):
        return template, version

    monkeypatch.setattr(readiness, "load_document_template_version", load)
    calls = []

    class Storage:
        async def read_persisted(self, **kwargs):
            calls.append(kwargs)
            return source_bytes(source_document("{{ customer.full_name }}"))

    result = await check_saved_document_readiness(
        None, tenant_scope=None, document=doc, template_storage=Storage()
    )
    assert not result["can_issue"] and fields(result) == {"customer.full_name"}
    assert calls[0]["version"] == 4 and calls[0]["storage_key"] == "frozen-key"
    assert frozen["meta"]["customer_readiness"] == metadata


@pytest.mark.asyncio
async def test_issued_document_readiness_never_reads_or_refreshes_existing_artifacts():
    result = await check_saved_document_readiness(
        None,
        tenant_scope=None,
        document=SimpleNamespace(status="issued", doc_type="contract"),
        template_storage=None,
    )
    assert result == {"checked": False, "missing_fields": [], "can_issue": True}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "status,guarded",
    [("draft", True), ("issued", False), ("sent", False), ("signed", False)],
)
async def test_docx_download_preserves_marker_after_external_edit_without_rewriting_storage(
    monkeypatch, status, guarded
):
    from modules.documents.application import artifact_helpers as helpers
    from modules.documents.application.lifecycle_service import ManagedDocumentService
    from modules.documents.application import customer_readiness as readiness

    document = SimpleNamespace(status=status, doc_type="contract")

    async def get_document(*args, **kwargs):
        return document

    monkeypatch.setattr(ManagedDocumentService, "get_document", get_document)

    async def check(*args, **kwargs):
        assert status == "draft"
        return {"can_issue": False}

    monkeypatch.setattr(readiness, "check_saved_document_readiness", check)
    artifact = SimpleNamespace(kind="source_docx", order_document_id=7)
    content = source_bytes(source_document("Внешняя правка без пометки"))
    assert (
        await helpers.requires_guarded_draft_download(
            None, tenant_scope=None, artifact=artifact
        )
        is guarded
    )
    output = await helpers.prepare_artifact_download(
        None,
        tenant_scope=None,
        artifact=artifact,
        content=content,
        template_storage=None,
    )
    if guarded:
        assert Document(BytesIO(output)).paragraphs[0].text == INCOMPLETE_DRAFT_MARKER
        assert (
            Document(BytesIO(content)).paragraphs[0].text
            == "Внешняя правка без пометки"
        )
    else:
        assert output is content


def test_old_snapshot_without_party_metadata_does_not_silently_exempt_signer_position():
    frozen = {"values": {}, "conditions": {}}
    result = snapshot_customer_readiness(
        frozen,
        source=source_bytes(source_document("{{ customer.signer_position }}")),
        document_type="contract",
    )
    assert not result["can_issue"] and fields(result) == {"customer.signer_position"}


@pytest.mark.asyncio
@pytest.mark.parametrize("edited", [False, True])
async def test_actual_draft_preview_reapplies_marker_to_native_and_unmarked_edited_source(
    monkeypatch, edited
):
    from modules.documents.application import draft_preview as preview
    from modules.documents.application.lifecycle_service import ManagedDocumentService

    frozen = snapshot()
    frozen["meta"]["customer_readiness"] = {"can_issue": True}
    original = deepcopy(frozen)
    document = SimpleNamespace(
        status="draft",
        doc_type="contract",
        render_snapshot=frozen,
        template_version_id=4,
        internal_reference="frozen-draft",
        id=7,
    )
    template = SimpleNamespace(id=1, tenant_id=3, doc_type="contract")
    version = SimpleNamespace(
        version=4,
        source_filename="contract.docx",
        source_storage_key="key",
        checksum_sha256="sha",
        placeholder_schema={"fields": ["customer.full_name"]},
    )

    async def get_document(*args, **kwargs):
        return document

    async def load_version(*args, **kwargs):
        return template, version

    async def artifacts(*args, **kwargs):
        return [SimpleNamespace(kind="source_docx")] if edited else []

    monkeypatch.setattr(ManagedDocumentService, "get_document", get_document)
    monkeypatch.setattr(ManagedDocumentService, "list_artifacts", artifacts)
    monkeypatch.setattr(preview, "load_document_template_version", load_version)
    monkeypatch.setattr(preview, "stored_artifact", lambda value: value)

    class Templates:
        async def read_persisted(self, **kwargs):
            return source_bytes(source_document("Клиент: {{ customer.full_name }}"))

    class Artifacts:
        async def read(self, stored):
            return source_bytes(source_document("Ручная правка без пометки"))

    class Pdf:
        def convert_docx(self, content, *, filename):
            output = Document(BytesIO(content))
            assert output.paragraphs[0].text == INCOMPLETE_DRAFT_MARKER
            assert (
                output.sections[0].header.paragraphs[0].text == INCOMPLETE_DRAFT_MARKER
            )
            assert ("Ручная правка" in output.paragraphs[1].text) is edited
            return b"%PDF-fake"

    pdf, name = await preview.ManagedDocumentDraftPreviewService.render_pdf(
        None,
        tenant_scope=SimpleNamespace(tenant_id=3),
        document_id=7,
        template_storage=Templates(),
        artifact_storage=Artifacts(),
        pdf_converter=Pdf(),
    )
    assert pdf == b"%PDF-fake" and name == "draft-frozen-draft.pdf"
    assert frozen == original


@pytest.mark.parametrize(
    "story",
    ["first_page_header", "first_page_footer", "even_page_header", "even_page_footer"],
)
@pytest.mark.parametrize("enabled", [False, True])
def test_disabled_word_story_fields_are_not_requirements_but_enabled_stories_are(
    story, enabled
):
    document = source_document("Клиент: {{ customer.full_name }}")
    section = document.sections[0]
    getattr(section, story).paragraphs[0].text = "{{ customer.acting_basis }}"
    if story.startswith("first"):
        section.different_first_page_header_footer = enabled
    else:
        document.settings.odd_and_even_pages_header_footer = enabled
    source = source_bytes(document)
    frozen = snapshot(**{"customer.full_name": "ООО Клиент"})
    result = snapshot_customer_readiness(
        frozen, source=source, document_type="contract"
    )
    assert result["can_issue"] is not enabled
    assert fields(result) == ({"customer.acting_basis"} if enabled else set())
    # Security/authoring discovery still inventories every Word story.
    assert "customer.acting_basis" in NativeDocxRenderer().discover_placeholders(source)


def test_later_enabled_section_includes_inherited_first_header_nested_table():
    document = source_document("Клиент: {{ customer.full_name }}")
    section = document.sections[0]
    section.different_first_page_header_footer = False
    table = section.first_page_header.add_table(
        rows=1, cols=1, width=section.page_width
    )
    table.cell(0, 0).add_table(rows=1, cols=1).cell(
        0, 0
    ).text = "{{ customer.acting_basis }}"
    document.add_section(
        WD_SECTION_START.NEW_PAGE
    ).different_first_page_header_footer = True
    frozen = snapshot(**{"customer.full_name": "ООО Клиент"})
    result = snapshot_customer_readiness(
        frozen, source=source_bytes(document), document_type="contract"
    )
    assert not result["can_issue"] and fields(result) == {"customer.acting_basis"}


@pytest.mark.parametrize(
    "frozen",
    [
        None,
        "bad",
        [],
        {},
        {"values": []},
        {"values": {"customer.full_name": {}}, "conditions": {}},
        {"values": {}, "conditions": []},
    ],
)
def test_malformed_factual_snapshot_has_a_readable_failure(frozen):
    with pytest.raises(ValueError, match="Сохранённый снимок документа повреждён"):
        snapshot_customer_readiness(
            frozen,
            source=source_bytes(source_document("{{ customer.full_name }}")),
            document_type="contract",
        )


def test_corrupt_persisted_docx_has_a_readable_failure():
    with pytest.raises(ValueError, match="Исходный DOCX шаблона повреждён"):
        snapshot_customer_readiness(
            snapshot(), source=b"corrupt-zip", document_type="contract"
        )


@pytest.mark.asyncio
@pytest.mark.parametrize("saved", [False, True])
@pytest.mark.parametrize(
    "error,status",
    [
        (OSError("storage unavailable"), 503),
        (FileNotFoundError("source missing"), 503),
        (TypeError("invalid context"), 400),
        (ValueError("corrupt source or snapshot"), 400),
    ],
)
async def test_readiness_http_endpoints_return_readable_errors_without_db_or_provider(
    monkeypatch, saved, error, status
):
    import importlib
    from fastapi import HTTPException
    from datetime import date
    from modules.documents.application.lifecycle_service import ManagedDocumentService
    from modules.documents.api.schemas import ManagedDocumentDraftPayload

    api = importlib.import_module("modules.documents.api.managed_document_readiness")
    compatibility = importlib.import_module("modules.documents.api.router")
    monkeypatch.setattr(
        compatibility,
        "get_private_attachment_storage",
        lambda: SimpleNamespace(provider_name="local", inventory_id="unit-private"),
    )

    async def fail(*args, **kwargs):
        raise error

    async def get(*args, **kwargs):
        return SimpleNamespace()

    monkeypatch.setattr(api, "check_selection_readiness", fail)
    monkeypatch.setattr(api, "check_saved_document_readiness", fail)
    monkeypatch.setattr(ManagedDocumentService, "get_document", get)
    auth = SimpleNamespace(tenant_scope=lambda: None)
    with pytest.raises(HTTPException) as caught:
        if saved:
            await api.get_managed_document_readiness(7, session=None, auth=auth)
        else:
            await api.check_managed_document_readiness(
                42,
                ManagedDocumentDraftPayload(
                    legal_entity_id=1,
                    document_type="invoice",
                    issue_date=date(2026, 10, 8),
                ),
                session=None,
                auth=auth,
            )
    assert caught.value.status_code == status
    assert caught.value.detail["message"]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "error,status",
    [
        (OSError("storage unavailable"), 503),
        (TypeError("malformed snapshot"), 409),
        (ValueError("corrupt source"), 409),
    ],
)
async def test_guarded_download_http_endpoint_handles_unavailable_corrupt_or_malformed_draft(
    monkeypatch, error, status
):
    import importlib
    from fastapi import HTTPException
    from modules.documents.application.lifecycle_service import ManagedDocumentService

    api = importlib.import_module("modules.documents.api.managed_documents_artifacts")
    artifact = SimpleNamespace(provider="local")

    async def get(*args, **kwargs):
        return artifact

    async def prepare(*args, **kwargs):
        raise error

    class Storage:
        async def read(self, value):
            return b"validated-stored-bytes"

    monkeypatch.setattr(ManagedDocumentService, "get_artifact", get)
    monkeypatch.setattr(ManagedDocumentService, "stored_artifact", lambda value: value)
    monkeypatch.setattr(
        api,
        "_legacy_private_storage",
        lambda *args: SimpleNamespace(provider_name="local", inventory_id="unit-private"),
    )
    monkeypatch.setattr(api, "PrivateDocumentArtifactStorage", lambda *args: Storage())
    monkeypatch.setattr(api, "prepare_artifact_download", prepare)
    with pytest.raises(HTTPException) as caught:
        await api.download_document_artifact(
            "artifact", session=None, auth=SimpleNamespace(tenant_scope=lambda: None)
        )
    assert caught.value.status_code == status and caught.value.detail["message"]
