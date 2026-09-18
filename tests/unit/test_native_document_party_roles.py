from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

from docx import Document
import pytest
from models import DocumentTemplate, DocumentTemplateVersion as StoredTemplateVersion

from modules.documents.application.context_builder import DocumentContextSelection
from modules.documents.application.party_roles import resolve_party_roles
from modules.documents.domain.roles import ROLE_FORMS
from modules.documents.infrastructure.renderers import DocumentTemplateVersion, NativeDocxRenderer, RenderContext
from modules.documents.infrastructure.renderers.docx_party_roles import detect_template_role_type


def source_bytes(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize('role_type', list(ROLE_FORMS))
def test_roles_inflect_and_preserve_formatting_without_replacing_customer_data(role_type):
    d = Document()
    p = d.add_paragraph('«')
    p.add_run('Про').bold = True
    p.add_run('давец').bold = True
    p.add_run('» и ПОКУПАТЕЛЬ. Продавца, продавцу, продавцом, продавце. ')
    p.add_run('{{ customer.full_name }}')
    d.add_table(rows=1, cols=1).cell(0, 0).text = 'Покупателя покупателю покупателем покупателе'
    d.sections[0].header.paragraphs[0].text = 'Исполнитель и Заказчик, субподрядчик'
    template = DocumentTemplateVersion(template_key='act', version=1, source=source_bytes(d), field_catalog=frozenset({'customer.full_name'}))
    output = NativeDocxRenderer().render(template, RenderContext(values={'customer.full_name': 'ООО «Продавец Покупатель»'}, document_role_type=role_type))
    result = Document(BytesIO(output.content))
    seller, buyer = ROLE_FORMS[role_type]
    assert result.paragraphs[0].text == (
        f'«{seller.nom.capitalize()}» и {buyer.nom.upper()}. {seller.gen.capitalize()}, {seller.dat}, {seller.ins}, {seller.prep}. ООО «Продавец Покупатель»'
    )
    assert ''.join(r.text for r in result.paragraphs[0].runs if r.bold) == seller.nom.capitalize()
    assert result.tables[0].cell(0, 0).text == f'{buyer.gen.capitalize()} {buyer.dat} {buyer.ins} {buyer.prep}'
    assert result.sections[0].header.paragraphs[0].text == f'{seller.nom.capitalize()} и {buyer.nom.capitalize()}, субподрядчик'


def test_existing_snapshots_without_role_leave_template_literals_unchanged():
    d = Document()
    d.add_paragraph('Исполнитель — Заказчик')
    template = DocumentTemplateVersion(template_key='act', version=1, source=source_bytes(d), field_catalog=frozenset())
    result = NativeDocxRenderer().render(template, RenderContext(values={}))
    assert Document(BytesIO(result.content)).paragraphs[0].text == 'Исполнитель — Заказчик'


@pytest.mark.parametrize('text,expected', [
    ('Продавец и Покупатель', 'seller_buyer'),
    ('Исполнителем и Заказчиком', 'executor_customer'),
    ('Подрядчик — Заказчик', 'contractor_customer'),
    ('Продавец Покупатель Исполнитель Заказчик', None),
    ('{{ seller.legal_name }} и {{ customer.full_name }}', None),
])
def test_template_role_detection_is_unambiguous(text, expected):
    d = Document()
    d.add_paragraph(text)
    assert detect_template_role_type(source_bytes(d)) == expected


def selection(role=None):
    return DocumentContextSelection(order_id=1, document_type='act', legal_entity_id=2, issue_date=date(2026, 9, 18), document_role_type=role)


def order(role='seller_buyer'):
    return SimpleNamespace(tenant_id=1, document_role_type=role, customer_contract=None)


@pytest.mark.asyncio
async def test_frozen_basis_beats_order_and_template_and_explicit_choice_wins():
    basis = SimpleNamespace(render_snapshot={'meta': {'document_role_type': 'contractor_customer'}})
    args = dict(order=order(), base_document=basis, base_contract=None, template=SimpleNamespace(document_role_type='executor_customer'))
    assert await resolve_party_roles(AsyncMock(), selection=selection(), **args) == ('contractor_customer', 'basis')
    assert await resolve_party_roles(AsyncMock(), selection=selection('executor_customer'), **args) == ('executor_customer', 'explicit')


@pytest.mark.asyncio
async def test_customer_contract_beats_conflicting_order_override_in_auto_mode():
    assert await resolve_party_roles(AsyncMock(), selection=selection(), order=order(), base_document=None, base_contract=SimpleNamespace(document_role_type='executor_customer')) == ('executor_customer', 'basis')


@pytest.mark.asyncio
async def test_historical_basis_uses_exact_version_not_current_template():
    d = Document()
    d.add_paragraph('Подрядчик — Заказчик')
    storage = SimpleNamespace(read_persisted=AsyncMock(return_value=source_bytes(d)))
    version = StoredTemplateVersion(id=7, template_id=5, version=3, renderer='docx', source_storage_key='source', source_filename='old.docx', checksum_sha256='checksum')
    template = DocumentTemplate(id=5, tenant_id=1, name='Old contract', doc_type='contract')
    session = SimpleNamespace(get=AsyncMock(side_effect=lambda model, _id: template if model is DocumentTemplate else version))
    basis = SimpleNamespace(render_snapshot={}, template_version_id=7, document_template_id=5)
    assert await resolve_party_roles(session, selection=selection(), order=order(), base_document=basis, base_contract=None, template_storage=storage) == ('contractor_customer', 'basis_template')
    assert storage.read_persisted.call_args.kwargs['version'] == 3
    assert storage.read_persisted.call_args.kwargs['tenant_id'] == 1
    template.tenant_id = 2
    with pytest.raises(ValueError, match='не принадлежит'):
        await resolve_party_roles(session, selection=selection(), order=order(), base_document=basis, base_contract=None, template_storage=storage)


@pytest.mark.asyncio
async def test_template_default_and_detected_roles_are_available_without_contract():
    assert await resolve_party_roles(AsyncMock(), selection=selection(), order=order(None), base_document=None, base_contract=None, template=SimpleNamespace(document_role_type='executor_customer')) == ('executor_customer', 'template')
    assert await resolve_party_roles(AsyncMock(), selection=selection(), order=order(None), base_document=None, base_contract=None) == ('seller_buyer', 'default')


@pytest.mark.asyncio
async def test_current_template_source_uses_order_tenant_with_real_version_model():
    document = Document()
    document.add_paragraph('Исполнитель и Заказчик')
    storage = SimpleNamespace(read_persisted=AsyncMock(return_value=source_bytes(document)))
    version = StoredTemplateVersion(
        id=7, template_id=5, version=3, renderer='docx',
        source_storage_key='source', source_filename='current.docx', checksum_sha256='checksum',
    )
    assert await resolve_party_roles(
        AsyncMock(), selection=selection(), order=order(None),
        base_document=None, base_contract=None, version=version, template_storage=storage,
    ) == ('executor_customer', 'template')
    assert storage.read_persisted.call_args.kwargs['tenant_id'] == 1
