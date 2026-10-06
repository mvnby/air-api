from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from modules.documents.infrastructure.renderers import (
    DocumentTemplateVersion, NativeDocxRenderer, RenderContext, TableBlockSpec,
)
from modules.documents.infrastructure.renderers.line_text import (
    LineTextRun, google_line_text_requests, parse_line_text,
)


def test_restricted_emphasis_keeps_newlines_and_non_markdown_text_literal():
    assert parse_line_text('**Работы:**\n*Проверка* <script>alert(1)</script> [ссылка](url) * без пары') == [
        LineTextRun('Работы:', bold=True),
        LineTextRun('\n'),
        LineTextRun('Проверка', italic=True),
        LineTextRun(' <script>alert(1)</script> [ссылка](url) * без пары'),
    ]
    assert parse_line_text('***Важно***') == [LineTextRun('Важно', bold=True, italic=True)]
    assert parse_line_text('**Работы *важно***') == [LineTextRun('Работы ', bold=True), LineTextRun('важно', bold=True, italic=True)]
    assert parse_line_text('*' * 10_000) == [LineTextRun('*' * 10_000)]
    assert parse_line_text('**Без пары*') == [LineTextRun('**Без пары*')]
    assert parse_line_text('Цена * 2 * 3') == [LineTextRun('Цена * 2 * 3')]


def test_docx_line_emphasis_preserves_template_style_split_fields_and_suffix():
    document = Document()
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text = '{{ lines }}{{ line.number }}'
    paragraph = cells[1].paragraphs[0]
    paragraph.add_run('До {{ line.').font.name = 'Arial'
    paragraph.add_run('title }} после {{ line.number }}').italic = True
    source = BytesIO()
    document.save(source)
    template = DocumentTemplateVersion(
        template_key='act', version=1, source=source.getvalue(), field_catalog=frozenset(),
        table_blocks=(TableBlockSpec(name='lines', row_fields=frozenset({'line.number', 'line.title'})),),
    )
    rendered = NativeDocxRenderer().render(template, RenderContext(values={}, table_rows={
        'lines': ({'line.number': '1', 'line.title': '**Работы:**\r\n*Осмотр* <b>literal</b>'},),
    }))
    paragraph = Document(BytesIO(rendered.content)).tables[0].cell(0, 1).paragraphs[0]
    assert paragraph.text == 'До Работы:\nОсмотр <b>literal</b> после 1'
    assert next(run for run in paragraph.runs if run.text == 'Работы:').bold is True
    assert next(run for run in paragraph.runs if run.text == 'Работы:').font.name == 'Arial'
    assert next(run for run in paragraph.runs if run.text == 'Осмотр').italic is True
    assert len(list(paragraph._p.iter(qn('w:br')))) == 1
    assert next(run for run in paragraph.runs if ' после 1' in run.text).italic is True


def test_google_emphasis_offsets_count_utf16_and_keep_html_literal():
    requests = google_line_text_requests('🔧 **Осмотр**\r\n*<b>текст</b>*', 12)
    assert requests[0]['insertText'] == {'location': {'index': 12}, 'text': '🔧 Осмотр\n<b>текст</b>'}
    assert requests[1]['updateTextStyle'] == {
        'range': {'startIndex': 15, 'endIndex': 21}, 'textStyle': {'bold': True}, 'fields': 'bold',
    }
    assert requests[2]['updateTextStyle']['range'] == {'startIndex': 22, 'endIndex': 34}
