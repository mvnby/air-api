from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlmodel import select

from models import (
    Customer,
    DocumentExternalEditSession,
    DocumentLegalEntity,
    DocumentTemplate,
    DocumentTemplateActLink,
    DocumentTemplateCustomerLink,
    DocumentTemplateVersion,
    StaffUser,
    Tenant,
)
from models.tenancy import TenantScope
from modules.documents.application.external_edit_sessions import (
    ExternalEditProviderError,
    ExternalEditSessionConflictError,
    ExternalEditSessionNotFoundError,
    TemplateExternalEditSessionService,
)
from modules.documents.application.template_versions import (
    NativeTemplatePlaceholderContract,
    NativeTemplateVersionService,
    TemplateVersionValidationError,
)
from modules.documents.infrastructure.external_edit_provider import (
    DOCX_CONTENT_TYPE,
    DownloadedExternalEditFile,
    ExternalEditFileMetadata,
)
from modules.documents.infrastructure.template_source_storage import (
    PrivateTemplateSourceStorage,
)
from services.private_attachment_storage_service import LocalPrivateAttachmentStorage


def _docx_bytes(*placeholders: str) -> bytes:
    document = Document()
    for placeholder in placeholders:
        document.add_paragraph(f"{{{{ {placeholder} }}}}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class _FakeExternalEditor:
    provider_name = "google_drive"

    def __init__(
        self,
        content: bytes,
        *,
        connection_id: str = "connection-1",
        file_id: str = "google-file-1",
    ) -> None:
        self.connection_id = connection_id
        self.file_id = file_id
        self.content = content
        self.revision = "revision-1"
        self.edit_session_id = ""
        self.ensure_calls = 0
        self.download_calls = 0
        self.metadata_calls = 0
        self.ensured_contents: list[bytes] = []

    def metadata(self) -> ExternalEditFileMetadata:
        return ExternalEditFileMetadata(
            file_id=self.file_id,
            edit_session_id=self.edit_session_id,
            edit_url="https://docs.google.com/document/d/google-file-1/edit",
            filename="Договор.docx",
            mime_type=DOCX_CONTENT_TYPE,
            revision=self.revision,
            modified_at=datetime(2026, 9, 2, tzinfo=timezone.utc),
        )

    async def ensure_docx(
        self, *, edit_session_id: str, filename: str, content: bytes
    ) -> ExternalEditFileMetadata:
        assert edit_session_id
        assert filename.endswith(".docx")
        assert content.startswith(b"PK")
        self.ensure_calls += 1
        self.ensured_contents.append(content)
        self.edit_session_id = edit_session_id
        return self.metadata()

    async def get_metadata(self, file_id: str) -> ExternalEditFileMetadata:
        assert file_id == self.file_id
        self.metadata_calls += 1
        return self.metadata()

    async def download_docx(self, file_id: str) -> DownloadedExternalEditFile:
        assert file_id == self.file_id
        self.download_calls += 1
        return DownloadedExternalEditFile(self.metadata(), self.content)


async def _database(tmp_path: Path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'external.db'}")
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        for model in (
            Tenant,
            DocumentLegalEntity,
            Customer,
            DocumentTemplate,
            DocumentTemplateCustomerLink,
            DocumentTemplateActLink,
            DocumentTemplateVersion,
            StaffUser,
            DocumentExternalEditSession,
        ):
            await connection.run_sync(model.__table__.create)
    return engine, sessions


async def _template_fixture(session, storage):
    tenant = Tenant(slug="owner", display_name="Owner")
    session.add(tenant)
    await session.commit()
    issuer = DocumentLegalEntity(
        tenant_id=int(tenant.id),
        slug="issuer",
        display_name="Issuer",
        requisites={},
    )
    session.add(issuer)
    await session.commit()
    scope = TenantScope(int(tenant.id), 1)
    template = await NativeTemplateVersionService.create_template(
        session,
        tenant_scope=scope,
        legal_entity_id=int(issuer.id),
        name="Договор",
        doc_type="contract",
    )
    content = _docx_bytes("document.official_number")
    version = await NativeTemplateVersionService.upload_native_docx_version(
        session,
        tenant_scope=scope,
        legal_entity_id=int(issuer.id),
        template_id=int(template.id),
        filename="Договор.docx",
        content=content,
        placeholder_contract=NativeTemplatePlaceholderContract.create(
            field_catalog={"document.official_number"}
        ),
        storage=storage,
    )
    await NativeTemplateVersionService.activate_version(
        session,
        tenant_scope=scope,
        legal_entity_id=int(issuer.id),
        template_id=int(template.id),
        version_id=int(version.id),
    )
    return scope, issuer, template, version, content


@pytest.mark.asyncio
async def test_google_template_session_is_idempotent_and_tenant_scoped(
    tmp_path: Path,
) -> None:
    engine, sessions = await _database(tmp_path)
    storage = PrivateTemplateSourceStorage(
        LocalPrivateAttachmentStorage(tmp_path / "private")
    )
    provider = _FakeExternalEditor(_docx_bytes("document.official_number"))
    try:
        async with sessions() as session:
            scope, issuer, template, version, _ = await _template_fixture(
                session, storage
            )
            first = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
            )
            second = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
            )

            assert first.id == second.id
            assert first.status == second.status == "ready"
            assert first.remote_revision == "revision-1"
            assert provider.ensure_calls == 1

            other_scope = TenantScope(scope.tenant_id + 100, 1)
            with pytest.raises(ExternalEditSessionNotFoundError):
                await TemplateExternalEditSessionService.get_session(
                    session,
                    tenant_scope=other_scope,
                    legal_entity_id=int(issuer.id),
                    template_id=int(template.id),
                    version_id=int(version.id),
                    provider=provider,
                )

            provider.edit_session_id = "another-session"
            with pytest.raises(ExternalEditProviderError):
                await TemplateExternalEditSessionService.get_session(
                    session,
                    tenant_scope=scope,
                    legal_entity_id=int(issuer.id),
                    template_id=int(template.id),
                    version_id=int(version.id),
                    provider=provider,
                )
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_status_refresh_does_not_break_an_active_sync_lease(
    tmp_path: Path,
) -> None:
    engine, sessions = await _database(tmp_path)
    storage = PrivateTemplateSourceStorage(
        LocalPrivateAttachmentStorage(tmp_path / "private")
    )
    provider = _FakeExternalEditor(_docx_bytes("document.official_number"))
    try:
        async with sessions() as session:
            scope, issuer, template, version, _ = await _template_fixture(
                session, storage
            )
            edit_session = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
            )
            edit_session.status = "syncing"
            edit_session.active_sync_key = "another-tab"
            session.add(edit_session)
            await session.commit()
            provider.metadata_calls = 0

            current = await TemplateExternalEditSessionService.get_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                provider=provider,
            )

            assert current.status == "syncing"
            assert current.active_sync_key == "another-tab"
            assert provider.metadata_calls == 0
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_remote_initialization_lease_prevents_duplicate_google_file(
    tmp_path: Path,
) -> None:
    engine, sessions = await _database(tmp_path)
    storage = PrivateTemplateSourceStorage(
        LocalPrivateAttachmentStorage(tmp_path / "private")
    )
    provider = _FakeExternalEditor(_docx_bytes("document.official_number"))
    try:
        async with sessions() as session:
            scope, issuer, template, version, _ = await _template_fixture(
                session, storage
            )
            edit_session = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
            )
            edit_session.remote_file_id = None
            edit_session.edit_url = None
            edit_session.status = "syncing"
            edit_session.active_sync_key = "init:another-worker"
            session.add(edit_session)
            await session.commit()

            with pytest.raises(
                ExternalEditSessionConflictError,
                match="другой вкладке",
            ):
                await TemplateExternalEditSessionService.ensure_session(
                    session,
                    tenant_scope=scope,
                    legal_entity_id=int(issuer.id),
                    template_id=int(template.id),
                    version_id=int(version.id),
                    source_storage=storage,
                    provider=provider,
                )
            assert provider.ensure_calls == 1
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_stale_remote_initialization_lease_can_recover(
    tmp_path: Path,
) -> None:
    engine, sessions = await _database(tmp_path)
    storage = PrivateTemplateSourceStorage(
        LocalPrivateAttachmentStorage(tmp_path / "private")
    )
    provider = _FakeExternalEditor(_docx_bytes("document.official_number"))
    try:
        async with sessions() as session:
            scope, issuer, template, version, _ = await _template_fixture(
                session, storage
            )
            edit_session = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
            )
            edit_session.remote_file_id = None
            edit_session.edit_url = None
            edit_session.status = "syncing"
            edit_session.active_sync_key = "init:abandoned-worker"
            edit_session.updated_at = datetime.now(timezone.utc) - timedelta(minutes=6)
            session.add(edit_session)
            await session.commit()

            recovered = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
            )
            assert recovered.status == "ready"
            assert recovered.active_sync_key is None
            assert provider.ensure_calls == 2
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_google_sync_creates_new_draft_without_overwriting_active_version(
    tmp_path: Path,
) -> None:
    engine, sessions = await _database(tmp_path)
    storage = PrivateTemplateSourceStorage(
        LocalPrivateAttachmentStorage(tmp_path / "private")
    )
    provider = _FakeExternalEditor(_docx_bytes("document.official_number"))
    try:
        async with sessions() as session:
            scope, issuer, template, version, _ = await _template_fixture(
                session, storage
            )
            staff_user = StaffUser(
                display_name="Редактор шаблонов",
                username="template-editor",
                roles=["manager"],
                primary_role="manager",
            )
            session.add(staff_user)
            await session.commit()
            edit_session = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
                staff_user_id=int(staff_user.id),
            )
            original_checksum = edit_session.base_checksum_sha256
            assert edit_session.created_by_staff_user_id == staff_user.id
            provider.content = _docx_bytes(
                "document.official_number", "customer.full_name"
            )
            provider.revision = "revision-2"
            changed = await TemplateExternalEditSessionService.get_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                provider=provider,
            )
            assert changed.status == "changed"

            result = await TemplateExternalEditSessionService.sync(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                expected_base_checksum_sha256=original_checksum,
                expected_remote_revision=changed.remote_revision,
                idempotency_key="sync-1",
                source_storage=storage,
                provider=provider,
                staff_user_id=int(staff_user.id),
            )

            assert result.edit_session.status == "ready"
            assert result.edit_session.last_synced_by_staff_user_id == staff_user.id
            assert result.new_template_version is not None
            assert result.new_template_version.version == 2
            assert result.new_template_version.status == "draft"
            assert result.new_template_version.placeholder_schema["fields"] == [
                "customer.full_name",
                "document.official_number",
            ]
            versions = (
                (
                    await session.execute(
                        select(DocumentTemplateVersion)
                        .where(DocumentTemplateVersion.template_id == template.id)
                        .order_by(DocumentTemplateVersion.version)
                    )
                )
                .scalars()
                .all()
            )
            assert [(item.version, item.status) for item in versions] == [
                (1, "active"),
                (2, "draft"),
            ]

            repeated = await TemplateExternalEditSessionService.sync(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                expected_base_checksum_sha256=original_checksum,
                expected_remote_revision=changed.remote_revision,
                idempotency_key="sync-1",
                source_storage=storage,
                provider=provider,
            )
            assert repeated.new_template_version.id == result.new_template_version.id
            assert provider.download_calls == 1

            with pytest.raises(
                ExternalEditSessionConflictError, match="другой версии файла"
            ):
                await TemplateExternalEditSessionService.sync(
                    session,
                    tenant_scope=scope,
                    legal_entity_id=int(issuer.id),
                    template_id=int(template.id),
                    version_id=int(version.id),
                    expected_base_checksum_sha256=result.edit_session.base_checksum_sha256,
                    expected_remote_revision=changed.remote_revision,
                    idempotency_key="sync-1",
                    source_storage=storage,
                    provider=provider,
                )
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_google_template_session_reconnect_creates_fresh_audited_binding(
    tmp_path: Path,
) -> None:
    engine, sessions = await _database(tmp_path)
    storage = PrivateTemplateSourceStorage(
        LocalPrivateAttachmentStorage(tmp_path / "private")
    )
    first_provider = _FakeExternalEditor(
        _docx_bytes("document.official_number"),
        connection_id="connection-1",
        file_id="google-file-1",
    )
    second_provider = _FakeExternalEditor(
        _docx_bytes("customer.full_name"),
        connection_id="connection-2",
        file_id="google-file-2",
    )
    try:
        async with sessions() as session:
            scope, issuer, template, version, original_content = await _template_fixture(
                session, storage
            )
            first = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=first_provider,
            )
            rebound = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=second_provider,
            )

            assert rebound.id != first.id
            assert rebound.template_version_id == first.template_version_id == version.id
            assert rebound.provider_connection_id == "connection-2"
            assert rebound.remote_file_id == "google-file-2"
            assert second_provider.ensure_calls == 1
            assert second_provider.ensured_contents == [original_content]
            rows = (
                await session.execute(
                    select(DocumentExternalEditSession).where(
                        DocumentExternalEditSession.template_version_id == version.id
                    )
                )
            ).scalars().all()
            assert {item.provider_connection_id for item in rows} == {
                "connection-1",
                "connection-2",
            }
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_google_sync_rejects_stale_remote_revision_and_invalid_docx(
    tmp_path: Path,
) -> None:
    engine, sessions = await _database(tmp_path)
    storage = PrivateTemplateSourceStorage(
        LocalPrivateAttachmentStorage(tmp_path / "private")
    )
    provider = _FakeExternalEditor(_docx_bytes("document.official_number"))
    try:
        async with sessions() as session:
            scope, issuer, template, version, _ = await _template_fixture(
                session, storage
            )
            persisted_template_id = int(template.id)
            edit_session = await TemplateExternalEditSessionService.ensure_session(
                session,
                tenant_scope=scope,
                legal_entity_id=int(issuer.id),
                template_id=int(template.id),
                version_id=int(version.id),
                source_storage=storage,
                provider=provider,
            )
            provider.content = _docx_bytes("internal.secret")
            provider.revision = "revision-3"

            with pytest.raises(ExternalEditSessionConflictError, match="изменился"):
                await TemplateExternalEditSessionService.sync(
                    session,
                    tenant_scope=scope,
                    legal_entity_id=int(issuer.id),
                    template_id=int(template.id),
                    version_id=int(version.id),
                    expected_base_checksum_sha256=edit_session.base_checksum_sha256,
                    expected_remote_revision="revision-2",
                    idempotency_key="stale-remote",
                    source_storage=storage,
                    provider=provider,
                )

            with pytest.raises(TemplateVersionValidationError):
                await TemplateExternalEditSessionService.sync(
                    session,
                    tenant_scope=scope,
                    legal_entity_id=int(issuer.id),
                    template_id=int(template.id),
                    version_id=int(version.id),
                    expected_base_checksum_sha256=edit_session.base_checksum_sha256,
                    expected_remote_revision="revision-3",
                    idempotency_key="invalid-docx",
                    source_storage=storage,
                    provider=provider,
                )

            versions = (
                (
                    await session.execute(
                        select(DocumentTemplateVersion).where(
                            DocumentTemplateVersion.template_id == persisted_template_id
                        )
                    )
                )
                .scalars()
                .all()
            )
            assert len(versions) == 1
    finally:
        await engine.dispose()
