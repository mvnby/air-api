from io import BytesIO

import pytest
from docx import Document

from services.email_contract_review_service import (
    ContractReviewInputError,
    EmailContractReviewService,
    ExtractedContract,
    _ProviderReport,
    extract_contract,
)


def test_docx_review_extracts_full_text_instead_of_intake_excerpt():
    document = Document()
    document.add_paragraph("Предмет договора. " + "Работы выполняются по заявке. " * 200)
    document.add_paragraph("Оплата производится в течение 90 дней после подписания акта.")
    content = BytesIO()
    document.save(content)

    extracted = extract_contract("agreement.docx", content.getvalue())

    assert len(extracted.text) > 4_000
    assert "Оплата производится в течение 90 дней" in extracted.text
    assert extracted.page_numbers_available is False


def test_review_keeps_only_quotes_present_on_cited_pdf_page():
    extracted = ExtractedContract(
        pages=(
            "Пункт 1. Срок оказания услуг составляет 10 дней.",
            "Пункт 2. Оплата производится через 90 дней после акта.",
        ),
        page_numbers_available=True,
    )
    report = _ProviderReport.model_validate({"risks": [
        {"topic": "Оплата", "clause": "2", "page": 2,
         "quote": "Оплата производится через 90 дней после акта.",
         "concern": "Долгая отсрочка", "proposal": "Согласовать аванс?"},
        {"topic": "Оплата", "clause": "2", "page": 1,
         "quote": "Оплата производится через 90 дней после акта.",
         "concern": "Неверная страница", "proposal": "Согласовать аванс?"},
        {"topic": "Штраф", "clause": "3", "page": 2,
         "quote": "Штраф 10%", "concern": "Выдумано", "proposal": "Снизить штраф?"},
    ]})

    risks = EmailContractReviewService._verified_risks(report, extracted)

    assert len(risks) == 1
    assert risks[0].page == 2
    assert risks[0].clause == "2"


def test_review_rejects_unsupported_or_too_long_document():
    with pytest.raises(ContractReviewInputError):
        extract_contract("agreement.xlsx", b"x" * 2000)
    with pytest.raises(ContractReviewInputError):
        extract_contract("agreement.docx", b"x" * (15 * 1024 * 1024 + 1))
