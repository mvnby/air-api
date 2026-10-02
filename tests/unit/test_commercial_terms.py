from decimal import Decimal

import pytest
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlmodel import SQLModel

from models import Order, Storefront, Tenant
from models.tenancy import TenantScope
from modules.documents.api.business_schemas import BusinessDocumentTermsPayload
from modules.documents.application.business_context import build_business_document_context
from modules.documents.domain.business_terms import BusinessDocumentTerms, PaymentScheduleItem
from schemas_commercial_terms import ManagerCommercialTermsUpdate
from services.commercial_terms_extraction import extract_commercial_terms, suggest_document_terms
from services.order_commercial_terms_service import CommercialTermsRevisionConflict, OrderCommercialTermsService, commercial_terms_response, store_source_terms


def test_unspecified_days_and_payment_trigger_are_not_invented():
    term = extract_commercial_terms("Оплата 100% в течение 50 дней после выполнения работ.", source="Письмо")[0]
    assert (term.due_days, term.day_kind, term.due_event) == (50, None, "after_work")
    assert term.evidence == "Оплата 100% в течение 50 дней после выполнения работ."
    assert suggest_document_terms([term]) is None
    unknown = extract_commercial_terms("Оплата 100% в течение 50 календарных дней с даты счёта.", source="Письмо")[0]
    assert unknown.due_event is None
    assert suggest_document_terms([unknown]) is None


@pytest.mark.parametrize("kind,expected", [("календарных", "calendar"), ("банковских", "banking"), ("рабочих", "working")])
def test_supported_days_preserved_in_document_suggestion(kind, expected):
    terms = extract_commercial_terms(f"Оплата 100% в течение 50 {kind} дней после выполнения работ.", source="Тендер")
    suggestion = suggest_document_terms(terms)
    assert suggestion.payment_schedule[0].due_day_kind == expected
    assert suggestion.payment_schedule[0].due_days == 50


def test_advance_is_not_assigned_an_arbitrary_event_or_remaining_stage():
    terms = extract_commercial_terms("Аванс 50%; Оплата 50% после подписания акта.", source="Письмо")
    assert terms[0].due_event is None
    assert suggest_document_terms(terms) is None
    valid = extract_commercial_terms("Аванс 50% до начала работ; Оплата 50% после подписания акта.", source="Письмо")
    assert len(suggest_document_terms(valid).payment_schedule) == 2
    not_whole = extract_commercial_terms("Аванс 40% до начала работ; Оплата 50% после подписания акта.", source="Письмо")
    assert suggest_document_terms(not_whole) is None


def test_delivery_basis_is_visible_and_date_is_not_payment_date():
    terms = extract_commercial_terms("Срок поставки в течение 10 рабочих дней после заключения договора.\nПоставка до 12.10.2026.", source="Тендер")
    assert terms[0].kind == "delivery"
    assert terms[0].day_kind == "working"
    assert "после заключения договора" in terms[0].trigger_text
    assert terms[1].deadline == "2026-10-12"


def test_extraction_preserves_reviewed_proposal_and_confirmation():
    order = Order(id=1, tenant_id=1, storefront_id=1, technical_meta={"commercial_terms": {
        "revision": 3, "confirmed": True, "proposed": {"payment_schedule": [{"share_percent": "100", "due_event": "before_supply", "due_day_kind": "banking"}]}}})
    store_source_terms(order, extract_commercial_terms("Оплата 100% в течение 50 рабочих дней после выполнения работ.", source="Новый документ"))
    response = commercial_terms_response(order)
    assert response.confirmed and response.revision == 3
    assert response.proposed.payment_schedule[0].due_event == "before_supply"
    assert response.customer_requested[0].due_event == "after_work"


def test_working_days_render_as_working_days():
    context = build_business_document_context(document_type="contract", terms=BusinessDocumentTerms(contract_scenario="services", payment_schedule=(PaymentScheduleItem(Decimal("100"), "after_work", 50, "working"),)), act_terms=None, order_additional_conditions=None, total_amount=Decimal("100"))
    assert "рабочих дней" in context.values["payment.summary"]


@pytest.mark.asyncio
async def test_tenant_scope_revision_conflict_and_caller_transaction(tmp_path, monkeypatch):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'terms.db'}")
    async with engine.begin() as connection:
        await connection.run_sync(SQLModel.metadata.create_all)
    try:
        async with AsyncSession(engine, expire_on_commit=False) as session:
            session.add_all([Tenant(id=710, slug="terms", display_name="Terms"), Storefront(id=711, tenant_id=710, slug="terms", display_name="Terms")])
            order = Order(tenant_id=710, storefront_id=711)
            session.add(order)
            await session.commit()
            order_id = order.id
            with pytest.raises(LookupError):
                await OrderCommercialTermsService.read(session, order_id=order_id, scope=TenantScope(tenant_id=712, storefront_id=713))
            await session.rollback()
            scope = TenantScope(tenant_id=710, storefront_id=711)
            proposal = BusinessDocumentTermsPayload(payment_schedule=[{"share_percent": 100, "due_event": "before_supply"}])
            payload = ManagerCommercialTermsUpdate(expected_revision=0, proposed=proposal, confirmed=True)
            async with session.begin():
                response = await OrderCommercialTermsService.update(session, order_id=order_id, scope=scope, payload=payload, username="manager")
                assert session.in_transaction() and response.revision == 1
            from io import BytesIO
            from docx import Document
            from unittest.mock import AsyncMock
            from services.email_contract_review_service import EmailContractReviewService
            document = Document()
            document.add_paragraph("Договор на выполнение монтажных работ. " * 8)
            document.add_paragraph("Оплата 100% в течение 50 рабочих дней после выполнения работ.")
            binary = BytesIO()
            document.save(binary)
            loader = AsyncMock(return_value=("договор.docx", binary.getvalue()))
            monkeypatch.setattr(EmailContractReviewService, "load_content", loader)
            extracted = await OrderCommercialTermsService.extract(session, order_id=order_id, scope=scope, attachment_ids=[50])
            assert extracted.customer_requested[0].source == "договор.docx"
            assert extracted.proposed.payment_schedule[0].due_event == "before_supply"
            assert extracted.confirmed and extracted.revision == 1
            with pytest.raises(LookupError):
                await OrderCommercialTermsService.extract(session, order_id=order_id, scope=TenantScope(tenant_id=712, storefront_id=713), attachment_ids=[50])
            assert loader.await_count == 1
            with pytest.raises(CommercialTermsRevisionConflict):
                await OrderCommercialTermsService.update(session, order_id=order_id, scope=scope, payload=payload, username="manager")
            with pytest.raises(PermissionError):
                await OrderCommercialTermsService.update(session, order_id=order_id, scope=TenantScope(tenant_id=710, storefront_id=711, demo_read_only=True), payload=payload, username="manager")
    finally:
        await engine.dispose()


def test_incomplete_source_and_contradictory_full_payment_schedules_block_suggestion():
    terms = extract_commercial_terms("Оплата 100% в течение 50 рабочих дней после выполнения работ.", source="Письмо")
    order = Order(id=4, tenant_id=1, storefront_id=1, technical_meta={"email_source_text_truncated": True})
    store_source_terms(order, terms)
    assert commercial_terms_response(order).suggested is None
    assert any("неполный" in warning for warning in commercial_terms_response(order).warnings)
    terms.extend(extract_commercial_terms("Оплата 100% до поставки.", source="Другой документ"))
    assert suggest_document_terms(terms) is None


def test_equipment_advance_without_trigger_remains_explicit_question():
    terms = extract_commercial_terms("100% предоплата оборудования.", source="Письмо")
    assert terms[0].share_percent == Decimal("100")
    assert terms[0].due_event is None
    assert suggest_document_terms(terms) is None


def test_long_delivery_basis_is_bounded_without_crashing_intake():
    raw = "Срок поставки в течение 10 рабочих дней после заключения договора " + "условие " * 120
    term = extract_commercial_terms(raw, source="Описание закупки")[0]
    assert len(term.evidence) > 500
    assert len(term.trigger_text) <= 500
    assert "после заключения договора" in term.trigger_text


def test_payment_stage_limit_returns_no_suggestion_instead_of_validation_error():
    terms = extract_commercial_terms("\n".join(["Оплата 4% после выполнения работ."] * 25), source="Тендер")
    assert len(terms) == 25
    assert suggest_document_terms(terms) is None


def test_term_limit_keeps_visible_excerpt_but_marks_omitted_contradictory_tail():
    raw = "Оплата 100% до поставки.\n" + "Срок поставки до 12.10.2026.\n" * 99 + "Оплата 100% после выполнения работ."
    terms = extract_commercial_terms(raw, source="Договор")
    assert len(terms) == 100
    assert any("больше 100" in issue for term in terms for issue in term.issues)
    assert suggest_document_terms(terms) is None


def test_overlong_relevant_clause_cannot_be_ignored_when_suggesting_prior_schedule():
    raw = "Оплата 100% до поставки.\n" + "Срок поставки в течение 10 рабочих дней после заключения договора " + "условие " * 400
    terms = extract_commercial_terms(raw, source="Договор")
    assert len(terms) == 2
    assert len(terms[1].evidence) == 2000
    assert suggest_document_terms(terms) is None
    assert any("не разобрано полностью" in issue for term in terms for issue in term.issues)


@pytest.mark.parametrize("meta", ["invalid", {"commercial_terms": "invalid"}, {"commercial_terms": {"customer_requested": {"broken": "shape"}}}, {"commercial_terms": {"customer_requested": [None, {"evidence": "без обязательных полей"}]}}, {"commercial_terms": {"proposed": {"payment_schedule": "broken"}, "confirmed": True, "revision": "not_integer"}}, {"belzakupki": []}])
def test_malformed_historical_metadata_does_not_break_inbox_projection(meta):
    from services.order_commercial_terms_service import commercial_terms_summary
    order = Order(id=10, tenant_id=1, storefront_id=1)
    order.technical_meta = meta
    assert commercial_terms_summary(order) == []
    response = commercial_terms_response(order)
    assert response.suggested is None and not response.confirmed


def test_invalid_saved_clause_does_not_make_remaining_full_schedule_look_complete():
    complete = extract_commercial_terms("Оплата 100% до поставки.", source="Договор")[0]
    order = Order(id=11, tenant_id=1, storefront_id=1, technical_meta={"commercial_terms": {"customer_requested": [complete.model_dump(mode="json"), {"kind": "payment", "evidence": "неполные данные"}]}})
    response = commercial_terms_response(order)
    assert response.customer_requested[0].evidence == complete.evidence
    assert response.suggested is None
    assert response.warnings


def test_document_defaults_never_silently_truncate_delivery_conditions():
    terms = extract_commercial_terms("Оплата 100% до поставки.\n" + "\n".join(["Срок поставки до 12.10.2026 " + "условие " * 120] * 10), source="Договор")
    assert len(terms) == 11
    assert suggest_document_terms(terms) is None
