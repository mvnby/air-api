"""Content controls must obey the same rendering/security rules as plain DOCX."""

from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from modules.documents.application.artifact_helpers import build_render_inputs
from modules.documents.application.native_template_discovery import (
    discover_native_placeholder_contract,
)
from modules.documents.infrastructure.renderers import (
    DocumentTemplateVersion,
    NativeDocxRenderer,
    RenderContext,
    TemplateValidationError,
)


def _save(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _wrap(*elements):
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    properties.append(OxmlElement("w:richText"))
    control.append(properties)
    content = OxmlElement("w:sdtContent")
    control.append(content)
    parent = elements[0].getparent()
    parent.insert(parent.index(elements[0]), control)
    for element in elements:
        content.append(element)
    return control


def _template(source):
    contract = discover_native_placeholder_contract(source)
    return DocumentTemplateVersion(
        template_key="test", version=1, source=source,
        field_catalog=contract.field_catalog,
        condition_catalog=contract.condition_catalog,
        table_blocks=contract.table_blocks,
    )


@pytest.mark.parametrize("location", [
    "body", "table", "row", "cell", "inline", "nested",
    "header", "footer", "first_page_header", "first_page_footer",
    "even_page_header", "even_page_footer",
])
def test_content_control_fields_are_discovered_and_rendered_with_formatting(location):
    document = Document()
    if location in {"body", "inline", "nested"}:
        paragraph = document.add_paragraph("Продавец: ")
    elif location in {"table", "row", "cell"}:
        table = document.add_table(rows=1, cols=1)
        paragraph = table.cell(0, 0).paragraphs[0]
    else:
        paragraph = getattr(document.sections[0], location).paragraphs[0]
    paragraph.add_run("{{ seller.").bold = True
    paragraph.add_run("legal_name }}")
    if location == "inline":
        _wrap(*[run._r for run in paragraph.runs[1:]])
    elif location in {"table", "row", "cell"}:
        _wrap({"table": table._tbl, "row": table.rows[0]._tr,
               "cell": table.cell(0, 0)._tc}[location])
    else:
        control = _wrap(paragraph._p)
        if location == "nested":
            _wrap(control)
    source = _save(document)
    original_xml = document._element.xml
    renderer = NativeDocxRenderer()
    assert renderer.discover_placeholder_counts(source) == {"seller.legal_name": 1}
    rendered = renderer.render(
        _template(source), RenderContext(values={"seller.legal_name": "ООО Продавец"})
    )
    output = Document(BytesIO(rendered.content))
    if location in {"table", "row", "cell"}:
        result = output.tables[0].cell(0, 0).paragraphs[0]
    elif location in {"body", "inline", "nested"}:
        result = output.paragraphs[0]
    else:
        result = getattr(output.sections[0], location).paragraphs[0]
    assert result.text.endswith("ООО Продавец")
    assert next(run for run in result.runs if "ООО" in run.text).bold is True
    assert not list(output._element.iter(qn("w:sdt")))
    assert document._element.xml == original_xml


@pytest.mark.parametrize("enabled", [True, False])
def test_conditions_inside_controls_filter_discovery_and_rendering(enabled):
    document = Document()
    start = document.add_paragraph("{{#if seller.is_vat_payer}}")
    body = document.add_paragraph("{{ seller.bank_name }}")
    end = document.add_paragraph("{{/if seller.is_vat_payer}}")
    _wrap(start._p, body._p, end._p)
    source = _save(document)
    renderer = NativeDocxRenderer()
    conditions = {"seller.is_vat_payer": enabled}
    assert renderer.discover_conditions(source) == {"seller.is_vat_payer"}
    assert renderer.discover_active_placeholders(source, conditions) == (
        {"seller.bank_name"} if enabled else set()
    )
    output = Document(BytesIO(renderer.render(
        _template(source),
        RenderContext(values={"seller.bank_name": "Банк"}, conditions=conditions),
    ).content))
    assert [p.text for p in output.paragraphs] == (["Банк"] if enabled else [])


@pytest.mark.parametrize("wrapper", ["table", "row", "cell"])
def test_repeatable_rows_inside_controls_keep_table_layout(wrapper):
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Оборудование"
    table.cell(1, 0).text = "{{ lines }}{{ line.title }}"
    _wrap({"table": table._tbl, "row": table.rows[1]._tr,
           "cell": table.cell(1, 0)._tc}[wrapper])
    source = _save(document)
    output = Document(BytesIO(NativeDocxRenderer().render(
        _template(source), RenderContext(values={}, table_rows={
            "lines": ({"line.title": "Первый"}, {"line.title": "Второй"}),
        }),
    ).content))
    assert [row.cells[0].text for row in output.tables[0].rows] == [
        "Оборудование", "Первый", "Второй",
    ]
    assert output.tables[0]._tbl.tblPr.xml == table._tbl.tblPr.xml


@pytest.mark.parametrize("placeholder", ["internal.secret", "seller.legal_name.upper()"])
def test_content_controls_cannot_hide_unknown_or_malformed_placeholders(placeholder):
    document = Document()
    _wrap(document.add_paragraph("{{ " + placeholder + " }}")._p)
    source = _save(document)
    with pytest.raises(TemplateValidationError):
        NativeDocxRenderer().render(_template(source), RenderContext(values={}))


def test_current_invoice_old_schema_recovers_seller_fields_without_changing_saved_data():
    document = Document()
    document.add_paragraph("Счёт № {{ document.official_full_number }}")
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].add_run("{{ seller.legal_name }}").bold = True
    cell.add_paragraph("{{ seller.bank_name }}")
    _wrap(table._tbl)
    document.add_paragraph("{{ seller.signer_name }}")
    source = _save(document)
    schema = {"fields": ["document.official_full_number", "seller.signer_name"]}
    snapshot = {"values": {
        "document.official_full_number": "С-123", "seller.legal_name": "ООО Продавец",
        "seller.bank_name": "Банк", "seller.signer_name": "Иванов",
        "internal.secret": "PRIVATE",
    }}
    original_schema, original_snapshot = deepcopy(schema), deepcopy(snapshot)
    template, context = build_render_inputs(
        template=SimpleNamespace(id=49, doc_type="invoice"),
        version=SimpleNamespace(version=3, source_filename="invoice.docx", placeholder_schema=schema),
        source=source, snapshot=snapshot,
    )
    output = Document(BytesIO(NativeDocxRenderer().render(template, context).content))
    assert output.tables[0].cell(0, 0).text == "ООО Продавец\nБанк"
    assert [p.text for p in output.paragraphs] == ["Счёт № С-123", "Иванов"]
    assert "internal.secret" not in context.values
    assert schema == original_schema
    assert snapshot == original_snapshot


def test_recovery_does_not_approve_ordinary_fields_omitted_from_explicit_catalogue():
    document = Document()
    document.add_paragraph("{{ seller.bank_name }}")
    _wrap(document.add_paragraph("{{ seller.legal_name }}")._p)
    template, context = build_render_inputs(
        template=SimpleNamespace(id=1),
        version=SimpleNamespace(version=1, source_filename="test.docx", placeholder_schema={}),
        source=_save(document), snapshot={"values": {"seller.legal_name": "Продавец"}},
    )
    with pytest.raises(TemplateValidationError, match="seller.bank_name"):
        NativeDocxRenderer().render(template, context)


def test_old_contract_recovers_conditions_and_table_blocks_from_controls():
    document = Document()
    start = document.add_paragraph("{{#if seller.is_vat_payer}}")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "{{ lines }}{{ line.title }}"
    end = document.add_paragraph("{{/if seller.is_vat_payer}}")
    _wrap(start._p, table._tbl, end._p)
    template, context = build_render_inputs(
        template=SimpleNamespace(id=1),
        version=SimpleNamespace(version=1, source_filename="test.docx", placeholder_schema={}),
        source=_save(document), snapshot={
            "values": {}, "conditions": {"seller.is_vat_payer": True},
            "table_rows": {"lines": [{"line.title": "Оборудование"}]},
        },
    )
    output = Document(BytesIO(NativeDocxRenderer().render(template, context).content))
    assert output.tables[0].cell(0, 0).text == "Оборудование"
    assert not output.paragraphs


def test_old_contract_recovery_cannot_approve_unknown_hidden_field():
    document = Document()
    _wrap(document.add_paragraph("{{ internal.secret }}")._p)
    template, context = build_render_inputs(
        template=SimpleNamespace(id=1),
        version=SimpleNamespace(version=1, source_filename="test.docx", placeholder_schema={}),
        source=_save(document), snapshot={"values": {"internal.secret": "PRIVATE"}},
    )
    with pytest.raises(TemplateValidationError, match="internal.secret"):
        NativeDocxRenderer().render(template, context)


def test_controls_cannot_hide_missing_customer_requirements_or_duplicate_official_markers():
    from modules.documents.application.customer_readiness import snapshot_customer_readiness
    from modules.documents.application.editable_draft import (
        EditableDraftError, official_placeholder_counts, validate_editable_draft,
    )

    document = Document()
    document.add_paragraph("{{ document.official_full_number }}")
    number = document.add_paragraph("{{ document.official_full_number }}")
    customer = document.add_paragraph("{{ customer.full_name }}")
    _wrap(number._p, customer._p)
    source = _save(document)
    readiness = snapshot_customer_readiness(
        {"values": {"customer.full_name": ""}}, source=source, document_type="invoice",
    )
    assert readiness["can_issue"] is False
    assert any(item["field"] == "customer.full_name" for item in readiness["missing_fields"])
    schema = {"fields": ["document.official_full_number"]}
    counts = official_placeholder_counts(source=source, placeholder_schema=schema)
    assert counts == {"document.official_full_number": 2}
    edited = Document()
    edited.add_paragraph("{{ document.official_full_number }}")
    with pytest.raises(EditableDraftError, match="изменено количество"):
        validate_editable_draft(
            source=_save(edited), placeholder_schema=schema, required_placeholder_counts=counts,
        )
