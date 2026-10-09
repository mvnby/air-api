from datetime import date
from io import BytesIO

import pytest
from docx import Document
from sqlmodel import select

from models import Customer, CustomerType, DocumentLegalEntity, DocumentNumberReservation, DocumentTemplate, Order, OrderStatus
from models.tenancy import TenantScope
from modules.documents.api.schemas import ManagedDocumentDraftPayload
from modules.documents.api.draft_selection import selection_from_payload
from modules.documents.application.context_builder import DocumentContextSelection
from modules.documents.application.lifecycle_service import ManagedDocumentService
from modules.documents.application.draft_preview import ManagedDocumentDraftPreviewService
from modules.documents.application.errors import ManagedDocumentConflictError, ManagedDocumentNotFoundError
from modules.documents.domain.participant_statement import ParticipantStatement
from modules.documents.infrastructure.artifact_storage import PrivateDocumentArtifactStorage
from modules.documents.infrastructure.template_source_storage import PrivateTemplateSourceStorage
from services.private_attachment_storage_service import LocalPrivateAttachmentStorage

FACTS = "Участник располагает оборудованием для выполнения работ.\nСрок выполнения — 30 календарных дней."


class RecordingPdfConverter:
    def __init__(self):
        self.contents = []

    def convert_docx(self, content, *, filename="document.docx"):
        self.contents.append(content)
        return b"%PDF-1.4\nparticipant-statement\n%%EOF"


async def seed(db, tmp_path, *, entity_type="organization"):
    customer = Customer(tenant_id=1, name="Заказчик закупки", inn="390000003", phone="+375291234567", type=CustomerType.company)
    issuer = DocumentLegalEntity(tenant_id=1, slug="statement-seller", display_name="Участник",
        legal_name="ИП Петров П.П." if entity_type == "individual_entrepreneur" else "ООО Участник",
        unp="390000002", entity_type=entity_type, is_default=True,
        requisites={"signer_name": "Петров П.П.", "signer_position": "Директор",
                    "acting_basis": "Устав", "signing_mode": "self" if entity_type == "individual_entrepreneur" else "statutory_body"})
    db.add_all([customer, issuer]); await db.flush()
    order = Order(tenant_id=1, storefront_id=1, customer_id=customer.id,
        status=OrderStatus.NEGOTIATION, title="Закупка из заявки",
        technical_meta={"procurement_reference": "2026-42"})
    db.add(order); await db.commit()
    private = LocalPrivateAttachmentStorage(tmp_path / "statement-storage")
    return TenantScope(tenant_id=1, storefront_id=1, is_system=True), order, issuer, PrivateTemplateSourceStorage(private), PrivateDocumentArtifactStorage(private)


def selection(order, issuer):
    payload = ManagedDocumentDraftPayload(legal_entity_id=issuer.id, document_type="participant_statement",
        issue_date=date(2026, 10, 9), participant_statement=dict(procedure_reference="https://buyer.test/2026-42",
            lot="1 — обслуживание", buyer_name="Заказчик закупки", declaration_text=FACTS))
    return selection_from_payload(order.id, payload)


@pytest.mark.asyncio
@pytest.mark.parametrize("entity_type", ["organization", "individual_entrepreneur"])
async def test_statement_first_use_preview_issue_frozen_facts(db, tmp_path, entity_type):
    scope, order, issuer, templates, artifacts = await seed(db, tmp_path, entity_type=entity_type)
    draft = await ManagedDocumentService.create_draft(db, tenant_scope=scope,
        selection=selection(order, issuer), template_storage=templates)
    assert draft.doc_type == "participant_statement"
    assert draft.render_snapshot["values"]["participant.declaration_text"] == FACTS
    assert draft.render_snapshot["table_rows"] == {}
    assert draft.official_number is None
    version = draft.template_version_id
    issuer.legal_name = "Новое название"; issuer.requisites = {"signer_name": "Другой"}
    db.add(issuer); await db.commit()
    converter = RecordingPdfConverter()
    pdf, _ = await ManagedDocumentDraftPreviewService.render_pdf(db, tenant_scope=scope, document_id=draft.id,
        template_storage=templates, artifact_storage=artifacts, pdf_converter=converter)
    assert pdf.startswith(b"%PDF")
    preview_text = "\n".join(p.text for p in Document(BytesIO(converter.contents[-1])).paragraphs)
    assert FACTS in preview_text and "Петров П.П." in preview_text
    assert "Новое название" not in preview_text and "Другой" not in preview_text
    if entity_type == "individual_entrepreneur":
        assert "Директор" not in preview_text and "Устав" not in preview_text
    kwargs = dict(tenant_scope=scope, document_id=draft.id, template_storage=templates,
        artifact_storage=artifacts, pdf_converter=converter)
    with pytest.raises(ManagedDocumentConflictError, match="подтвердите"):
        await ManagedDocumentService.issue(db, **kwargs)
    assert await db.scalar(select(DocumentNumberReservation.id)) is None
    result = await ManagedDocumentService.issue(db, **kwargs,
        participant_statement_confirmed=True, participant_statement_confirmed_by="manager")
    assert result.document.status == "issued"
    assert result.document.template_version_id == version
    assert {a.kind for a in result.artifacts} == {"rendered_docx", "pdf"}
    text = "\n".join(p.text for p in Document(BytesIO(converter.contents[-1])).paragraphs)
    assert FACTS in text and "2026-42" in text and "1 — обслуживание" in text
    assert "{{" not in text and "судим" not in text and "задолж" not in text
    confirmation = result.document.render_snapshot["meta"]["participant_statement_confirmation"]
    assert confirmation["confirmed_by"] == "manager" and confirmation["issued_docx_checksum"]
    again = await ManagedDocumentService.issue(db, **kwargs)
    assert again.document.official_number == result.document.official_number


@pytest.mark.asyncio
async def test_statement_tenant_closed_and_disabled_template_guards(db, tmp_path):
    scope, order, issuer, templates, artifacts = await seed(db, tmp_path)
    with pytest.raises(ManagedDocumentNotFoundError):
        await ManagedDocumentService.create_draft(db, tenant_scope=TenantScope(tenant_id=99, storefront_id=99),
            selection=selection(order, issuer), template_storage=templates)
    draft = await ManagedDocumentService.create_draft(db, tenant_scope=scope,
        selection=selection(order, issuer), template_storage=templates)
    template = await db.get(DocumentTemplate, draft.document_template_id)
    template.is_active = False; db.add(template); await db.commit()
    with pytest.raises(ManagedDocumentConflictError, match="отключён"):
        await ManagedDocumentService.create_draft(db, tenant_scope=scope,
            selection=selection(order, issuer), template_storage=templates)
    await db.rollback()
    order.status = OrderStatus.CLOSED; db.add(order); await db.commit()
    with pytest.raises(ManagedDocumentConflictError, match="завершён"):
        await ManagedDocumentService.issue(db, tenant_scope=scope, document_id=draft.id,
            template_storage=templates, artifact_storage=artifacts, pdf_converter=RecordingPdfConverter(),
            participant_statement_confirmed=True)


def test_statement_schema_requires_explicit_text_and_rejects_wrong_document_type():
    from pydantic import ValidationError
    kwargs = dict(legal_entity_id=1, issue_date=date(2026, 10, 9))
    with pytest.raises(ValidationError):
        ManagedDocumentDraftPayload(**kwargs, document_type="participant_statement")
    with pytest.raises(ValidationError):
        ManagedDocumentDraftPayload(**kwargs, document_type="offer", participant_statement=dict(
            procedure_reference="42", lot="1", buyer_name="Заказчик", declaration_text=FACTS))
    with pytest.raises(ValueError):
        ParticipantStatement("42", "1", "Заказчик", "  ").values()


@pytest.mark.asyncio
async def test_statement_manual_text_edit_is_frozen_in_issued_artifacts(db, tmp_path):
    from modules.documents.application.managed_document_external_edits import ManagedDocumentExternalEditSessionService as Edits
    from modules.documents.application.editable_draft_issue import verify_document_external_edit_before_issue
    from tests.integration.test_managed_document_google_editing import FakeGoogleEditor

    scope, order, issuer, templates, artifacts = await seed(db, tmp_path)
    draft = await ManagedDocumentService.create_draft(db, tenant_scope=scope,
        selection=selection(order, issuer), template_storage=templates)
    provider = FakeGoogleEditor()
    edit = await Edits.ensure_session(db, tenant_scope=scope, document_id=draft.id,
        template_storage=templates, artifact_storage=artifacts, provider=provider)
    provider.add_manual_paragraph("Уточнение согласно форме заказчика: срок 45 дней.")
    changed = await Edits.get_session(db, tenant_scope=scope, document_id=draft.id, provider=provider)
    await Edits.sync(db, tenant_scope=scope, document_id=draft.id,
        expected_base_checksum_sha256=edit.base_checksum_sha256,
        expected_remote_revision=changed.remote_revision, idempotency_key="statement-sync",
        artifact_storage=artifacts, provider=provider)
    revision = await verify_document_external_edit_before_issue(db, tenant_scope=scope,
        document_id=draft.id, provider=provider)
    converter = RecordingPdfConverter()
    kwargs = dict(tenant_scope=scope, document_id=draft.id, template_storage=templates,
        artifact_storage=artifacts, pdf_converter=converter, verified_remote_revision=revision)
    with pytest.raises(ManagedDocumentConflictError, match="подтвердите"):
        await ManagedDocumentService.issue(db, **kwargs)
    result = await ManagedDocumentService.issue(db, **kwargs, participant_statement_confirmed=True,
        participant_statement_confirmed_by="manager")
    text = "\n".join(p.text for p in Document(BytesIO(converter.contents[-1])).paragraphs)
    assert "Уточнение согласно форме заказчика: срок 45 дней." in text
    assert "{{" not in text
    confirmation = result.document.render_snapshot["meta"]["participant_statement_confirmation"]
    assert confirmation["source_docx_checksum"] and confirmation["issued_docx_checksum"]
    provider.add_manual_paragraph("Не должно изменить выпущенную версию")
    rendered = next(a for a in result.artifacts if a.kind == "rendered_docx")
    persisted = await artifacts.read(ManagedDocumentService.stored_artifact(rendered))
    persisted_text = "\n".join(p.text for p in Document(BytesIO(persisted)).paragraphs)
    assert "Не должно изменить" not in persisted_text
