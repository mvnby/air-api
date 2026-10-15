from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from docx import Document

import pytest
from sqlalchemy import func
from sqlmodel import select

from core.security import create_access_token
from models import DocumentLegalEntity, DocumentTemplate, DocumentTemplateVersion, OrderDocument, Product, StaffUser, Storefront, StorefrontDomain, Supplier, Tenant, TenantCatalogGrant, TenantMembership, TenantOffer
from models.supplier import ProductSupplierMapping, SupplierOffer
from models.tenancy import TenantScope
from services.catalog_decision_projection import CatalogDecisionFilters, CatalogDecisionQueryService
from services.manager_tenant_catalog_reader import ManagerTenantCatalogReader
from models.communication import CommunicationDelivery
from models.integration_event import IntegrationOutboxEvent
from services.tenant_training_bootstrap_service import (
    SLUG, TenantTrainingBootstrapService, TenantTrainingBootstrapToken,
)
from services.tenant_training_setup_service import TenantTrainingSetupBlockedError


async def products(db):
    values = [Product(title=f"Complete split {i}", slug=f"bootstrap-complete-split-{i}",
                      price=2000 + i * 500, product_kind="complete_split_system") for i in range(2)]
    db.add_all(values)
    await db.commit()
    return values


@pytest.mark.asyncio
async def test_bootstrap_creates_private_scoped_target_atomically_and_retries(db):
    chosen = await products(db)
    ids = [p.id for p in chosen]
    plan = await TenantTrainingBootstrapService.plan(db, product_ids=ids)
    assert plan["status"] == "planned" and not plan["blockers"]
    result = await TenantTrainingBootstrapService.execute(db, product_ids=ids,
        plan_token=TenantTrainingBootstrapToken.issue(plan_digest=plan["plan_digest"]))
    await db.commit()
    target = result["target"]
    tenant = await db.get(Tenant, target["tenant_id"])
    storefront = await db.get(Storefront, target["storefront_id"])
    issuer = await db.get(DocumentLegalEntity, target["legal_entity_id"])
    assert tenant.slug == SLUG and not tenant.demo_read_only and not tenant.is_system
    assert storefront.status == "active" and storefront.is_default
    assert issuer.tenant_id == tenant.id and issuer.unp is None and issuer.is_default
    assert not any(issuer.requisites.get(k) for k in ("iban", "bank_name", "bic", "phone", "signature", "seal"))
    offers = list((await db.execute(select(TenantOffer).where(TenantOffer.tenant_id == tenant.id))).scalars())
    assert {o.product_id for o in offers} == set(ids)
    grant = await db.get(TenantCatalogGrant, target["catalog_grant_id"])
    assert grant.owner_type == "system" and grant.status == "active" and grant.mode == "all_published"
    assert grant.price_policy == "inherit_master" and grant.last_completed_sync_at is None
    assert all(o.catalog_grant_id == grant.id and o.price > 0 and o.price_source == "inherited_master" for o in offers)
    scope = TenantScope(tenant_id=tenant.id, storefront_id=storefront.id,
                        is_system=False, demo_read_only=False, is_canonical_storefront=False)
    catalog = await CatalogDecisionQueryService.list_products(db, tenant_scope=scope,
        filters=CatalogDecisionFilters(), page=1, limit=100, sort="title", direction="asc")
    assert {item["id"] for item in catalog["items"]} == set(ids)
    for product_id in ids:
        detail = await ManagerTenantCatalogReader.get_product(db, product_id=product_id, tenant_scope=scope)
        assert detail["allowed"] is True and detail["selection_item"]["id"] == product_id
    for model in (StorefrontDomain, CommunicationDelivery, IntegrationOutboxEvent):
        assert (await db.execute(select(func.count()).select_from(model))).scalar_one() == 0
    retry_plan = await TenantTrainingBootstrapService.plan(db, product_ids=ids)
    assert retry_plan["status"] == "already_created" and not retry_plan["blockers"]
    retry = await TenantTrainingBootstrapService.execute(db, product_ids=ids,
        plan_token=TenantTrainingBootstrapToken.issue(plan_digest=retry_plan["plan_digest"]))
    assert retry == {"status": "already_created", "target": target}


@pytest.mark.asyncio
async def test_existing_unowned_training_slug_is_never_adopted(db):
    chosen = await products(db)
    db.add(Tenant(slug=SLUG, display_name="Existing cabinet"))
    await db.commit()
    plan = await TenantTrainingBootstrapService.plan(db, product_ids=[p.id for p in chosen])
    assert plan["status"] == "blocked"
    assert any("never adopt" in message for message in plan["blockers"])


@pytest.mark.asyncio
async def test_changed_canonical_retail_blocks_before_tenant_creation(db):
    chosen = await products(db)
    ids = [p.id for p in chosen]
    plan = await TenantTrainingBootstrapService.plan(db, product_ids=ids)
    chosen[0].price += 1
    await db.commit()
    with pytest.raises(TenantTrainingSetupBlockedError, match="State changed"):
        await TenantTrainingBootstrapService.execute(db, product_ids=ids,
            plan_token=TenantTrainingBootstrapToken.issue(plan_digest=plan["plan_digest"]))
    assert (await db.execute(select(Tenant).where(Tenant.slug == SLUG))).first() is None


@pytest.mark.asyncio
async def test_partial_block_is_not_a_complete_split(db):
    chosen = await products(db)
    chosen[1].product_kind = "indoor_unit"
    await db.commit()
    plan = await TenantTrainingBootstrapService.plan(db, product_ids=[p.id for p in chosen])
    assert plan["status"] == "blocked"
    assert any("complete split systems" in message for message in plan["blockers"])


@pytest.mark.asyncio
@pytest.mark.parametrize("specs", [
    {"type": "внутренний блок"},
    {"type": "мульти-сплит-система"},
    {"indoor_units_count": "2"},
])
async def test_contradictory_source_metadata_blocks_bootstrap_before_any_write(db, specs):
    chosen = await products(db)
    ids = [p.id for p in chosen]
    reviewed = await TenantTrainingBootstrapService.plan(db, product_ids=ids)
    chosen[1].specs = specs
    await db.commit()
    plan = await TenantTrainingBootstrapService.plan(db, product_ids=ids)
    assert chosen[1].product_kind == "complete_split_system"
    assert plan["status"] == "blocked"
    assert plan["selection_eligible_product_ids"] == [ids[0]]
    assert "both reviewed products must pass current complete-split selection admission" in plan["blockers"]
    with pytest.raises(TenantTrainingSetupBlockedError, match="State changed"):
        await TenantTrainingBootstrapService.execute(db, product_ids=ids,
            plan_token=TenantTrainingBootstrapToken.issue(plan_digest=reviewed["plan_digest"]))
    assert (await db.execute(select(Tenant).where(Tenant.slug == SLUG))).first() is None


@pytest.mark.asyncio
async def test_private_recording_masks_supplier_costs_and_keeps_owner_editing(db, async_client):
    chosen = await products(db)
    supplier = Supplier(name="Actual supplier", code="training-cost-source", is_active=True)
    db.add(supplier)
    await db.flush()
    db.add_all([
        SupplierOffer(supplier_id=supplier.id, external_id="TRAINING-COST", wholesale_value=400,
                      wholesale_currency="BYN", rrc_byn=3000, qty=10),
        ProductSupplierMapping(product_id=chosen[0].id, supplier_id=supplier.id, external_id="TRAINING-COST"),
    ])
    await db.commit()
    ids = [p.id for p in chosen]
    plan = await TenantTrainingBootstrapService.plan(db, product_ids=ids)
    result = await TenantTrainingBootstrapService.execute(db, product_ids=ids,
        plan_token=TenantTrainingBootstrapToken.issue(plan_digest=plan["plan_digest"]))
    target = result["target"]
    owner = StaffUser(display_name="Training owner", username="training-owner-test", status="active",
                      primary_role="owner", roles=["owner"])
    db.add(owner)
    await db.flush()
    db.add(TenantMembership(tenant_id=target["tenant_id"], staff_user_id=owner.id, role="owner", status="active"))
    await db.commit()
    scope = TenantScope(tenant_id=target["tenant_id"], storefront_id=target["storefront_id"],
                        is_system=False, demo_read_only=False, is_canonical_storefront=False)
    detail = await ManagerTenantCatalogReader.get_product(db, product_id=ids[0], tenant_scope=scope)
    assert detail["selection_item"]["purchase_cost_byn"] == 2700  # 90% of RRC, never the real 400.
    snapshots = await CatalogDecisionQueryService.get_product_snapshots(db, tenant_scope=scope, product_ids=ids)
    assert snapshots[ids[0]].purchase_cost_byn == 2700
    token = create_access_token({"sub": owner.username, "staff_user_id": owner.id,
        "auth_version": owner.auth_version, "auth_source": "training-test"}, expires_delta=timedelta(minutes=5))
    response = await async_client.patch(f"/api/manager/orders/{result['records']['order_ids'][0]}",
        json={"title": "[ОБУЧЕНИЕ] Изменено во время урока"}, headers={"Authorization": f"Bearer {token}"})
    assert response.status_code == 200, response.text
    assert response.json()["title"] == "[ОБУЧЕНИЕ] Изменено во время урока"
    unrelated = Tenant(slug="training-kitlane-extra", display_name="Other partner")
    db.add(unrelated)
    await db.flush()
    other_scope = TenantScope(tenant_id=unrelated.id, storefront_id=scope.storefront_id,
        is_system=False, demo_read_only=False, is_canonical_storefront=False)
    assert await CatalogDecisionQueryService._uses_demo_cost(db, other_scope) is False


@pytest.mark.asyncio
async def test_training_native_bundle_import_and_read_only_preflight_need_no_real_requisites(db, tmp_path):
    from modules.documents.application.artifact_helpers import build_render_inputs
    from modules.documents.application.context_builder import DocumentContextBuilder, DocumentContextSelection
    from modules.documents.application.customer_readiness import check_selection_readiness
    from modules.documents.domain import BusinessDocumentTerms, PaymentScheduleItem
    from modules.documents.infrastructure.renderers import NativeDocxRenderer
    from modules.documents.infrastructure.template_source_storage import PrivateTemplateSourceStorage
    from scripts.manage_native_document_template_bundle import _apply_bundle_item, load_bundle, prepare_sources
    from scripts.prepare_training_document_templates import prepare
    from services.private_attachment_storage_service import LocalPrivateAttachmentStorage

    chosen = await products(db)
    ids = [p.id for p in chosen]
    plan = await TenantTrainingBootstrapService.plan(db, product_ids=ids)
    result = await TenantTrainingBootstrapService.execute(db, product_ids=ids,
        plan_token=TenantTrainingBootstrapToken.issue(plan_digest=plan["plan_digest"]))
    await db.commit()
    target = result["target"]
    scope = TenantScope(tenant_id=target["tenant_id"], storefront_id=target["storefront_id"],
                        is_system=False, demo_read_only=False, is_canonical_storefront=False)
    source_dir = tmp_path / "sources"
    prepare(source_dir)
    bundle = load_bundle(source_dir / "bundle.json")
    sources = prepare_sources(bundle, source_dir)
    storage = PrivateTemplateSourceStorage(LocalPrivateAttachmentStorage(tmp_path / "private"))
    for spec in bundle.templates:
        applied = await _apply_bundle_item(db, scope=scope, legal_entity_id=target["legal_entity_id"],
            bundle=bundle, spec=spec, source=sources[spec.key], storage=storage)
        selection = DocumentContextSelection(order_id=result["records"]["order_ids"][1],
            document_type=spec.doc_type, legal_entity_id=target["legal_entity_id"], issue_date=date.today(),
            business_terms=BusinessDocumentTerms(contract_scenario="supply",
                payment_schedule=(PaymentScheduleItem(Decimal("100"), "before_supply"),))
            if spec.doc_type == "contract" else None)
        readiness = await check_selection_readiness(db, tenant_scope=scope,
            selection=selection, template_id=applied["template_id"], template_storage=storage)
        assert readiness["can_issue"] is True and readiness["missing_fields"] == []
        template = await db.get(DocumentTemplate, applied["template_id"])
        version = await db.get(DocumentTemplateVersion, applied["activated_version_id"])
        snapshot = await DocumentContextBuilder.build(db, tenant_scope=scope, selection=selection,
            template=template, template_version=version, template_storage=storage)
        render_template, context = build_render_inputs(template=template, version=version,
            source=sources[spec.key].content, snapshot=snapshot)
        rendered = NativeDocxRenderer().render(render_template, context)
        document = Document(BytesIO(rendered.content))
        assert len(document.tables[0].rows) == 2  # Header plus the selected proposal's one row.
        assert chosen[0].title in document.tables[0].cell(1, 0).text
        assert chosen[1].title not in document.tables[0].cell(1, 0).text
        assert any("УЧЕБНЫЙ ОБРАЗЕЦ" in p.text for p in document.paragraphs)
    assert (await db.execute(select(func.count()).select_from(OrderDocument))).scalar_one() == 0
