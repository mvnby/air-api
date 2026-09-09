"""Regression coverage for staged B2C DOCX issuance and editable drafts."""

from datetime import date
from io import BytesIO

import pytest
from docx import Document
from sqlmodel import select

from models import OrderProductLink
from models.tenancy import TenantScope
from modules.documents.application import (
    DocumentContextSelection,
    ManagedDocumentService,
    NativeTemplateVersionService,
)
from modules.documents.application.editable_draft_issue import (
    verify_document_external_edit_before_issue,
)
from modules.documents.application.managed_document_external_edits import (
    ManagedDocumentExternalEditSessionService,
)
from modules.documents.application.native_template_discovery import (
    discover_native_placeholder_contract,
)
from modules.documents.domain import ConsumerDocumentTerms
from modules.documents.infrastructure.artifact_storage import PrivateDocumentArtifactStorage
from modules.documents.infrastructure.template_source_storage import PrivateTemplateSourceStorage
from services.private_attachment_storage_service import LocalPrivateAttachmentStorage
from tests.integration.test_document_context_builder import _seed_order
from tests.integration.test_managed_document_google_editing import (
    FakeGoogleEditor,
    _document_text,
)
from tests.integration.test_managed_document_lifecycle import FakePdfConverter


def _staged_template() -> bytes:
    document = Document()
    document.add_paragraph("{{#if installation.two_stages}}")
    document.add_paragraph(
        "ПРИНЯТ ТОЛЬКО ПЕРВЫЙ ЭТАП. "
        "Первый: {{ installation.first_stage_works }} "
        "Второй: {{ installation.second_stage_works }}"
    )
    document.add_paragraph("Первый этап: {{ installation.first_stage_amount }}")
    document.add_paragraph("Остаток: {{ installation.remaining_amount }}")
    document.add_paragraph("Срок второго этапа: {{ installation.second_stage_due }}")
    document.add_paragraph("{{/if installation.two_stages}}")
    document.add_paragraph("{{#if installation.single_stage}}")
    document.add_paragraph("СТАРОЕ ПОЛНОЕ ПРИНЯТИЕ ВСЕХ РАБОТ")
    document.add_paragraph("{{/if installation.single_stage}}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


async def _staged_fixture(db, tmp_path):
    order, issuer, proposal, _ = await _seed_order(db)
    issuer.requisites = {
        **issuer.requisites,
        "offer_url": "https://example.test/offer",
        "offer_version": "1",
        "offer_published_on": "09.09.2026",
    }
    link = (
        await db.execute(
            select(OrderProductLink).where(OrderProductLink.proposal_id == proposal.id)
        )
    ).scalar_one()
    link.price = 2740  # Together with two 200-BYN service lines: 3140 BYN.
    db.add_all((issuer, link))
    await db.commit()

    scope = TenantScope(tenant_id=1, storefront_id=1, is_system=True)
    private = LocalPrivateAttachmentStorage(tmp_path / "staged-installation-docx")
    template_storage = PrivateTemplateSourceStorage(private)
    artifact_storage = PrivateDocumentArtifactStorage(private)
    template = await NativeTemplateVersionService.create_template(
        db,
        tenant_scope=scope,
        legal_entity_id=issuer.id,
        name="Заказ-акт: поставка и монтаж",
        doc_type="b2c_supply_installation_act",
    )
    source = _staged_template()
    version = await NativeTemplateVersionService.upload_native_docx_version(
        db,
        tenant_scope=scope,
        legal_entity_id=issuer.id,
        template_id=template.id,
        filename="staged.docx",
        content=source,
        placeholder_contract=discover_native_placeholder_contract(source),
        storage=template_storage,
    )
    await NativeTemplateVersionService.activate_version(
        db,
        tenant_scope=scope,
        legal_entity_id=issuer.id,
        template_id=template.id,
        version_id=version.id,
    )
    return scope, order, issuer, proposal, template, template_storage, artifact_storage


@pytest.mark.asyncio
@pytest.mark.parametrize("editable", (False, True), ids=("native", "google-editable"))
@pytest.mark.parametrize("outdoor_in_first", (True, False), ids=("outdoor-first", "outdoor-second"))
async def test_staged_installation_issue_keeps_first_stage_acceptance_and_unit_assignment(
    db, tmp_path, editable, outdoor_in_first
):
    (
        scope,
        order,
        issuer,
        proposal,
        template,
        template_storage,
        artifact_storage,
    ) = await _staged_fixture(db, tmp_path)
    draft = await ManagedDocumentService.create_draft(
        db,
        tenant_scope=scope,
        template_id=template.id,
        selection=DocumentContextSelection(
            order_id=order.id,
            legal_entity_id=issuer.id,
            document_type="b2c_supply_installation_act",
            issue_date=date(2026, 9, 9),
            proposal_id=proposal.id,
            consumer_terms=ConsumerDocumentTerms(
                installation_two_stages=True,
                installation_first_stage_amount="3000",
                installation_outdoor_unit_in_first_stage=outdoor_in_first,
            ),
        ),
    )

    verified_revision = None
    if editable:
        provider = FakeGoogleEditor()
        session = await ManagedDocumentExternalEditSessionService.ensure_session(
            db,
            tenant_scope=scope,
            document_id=draft.id,
            template_storage=template_storage,
            artifact_storage=artifact_storage,
            provider=provider,
        )
        editable_text = _document_text(provider.content)
        assert "ПРИНЯТ ТОЛЬКО ПЕРВЫЙ ЭТАП" in editable_text
        assert "СТАРОЕ ПОЛНОЕ ПРИНЯТИЕ" not in editable_text
        verified_revision = await verify_document_external_edit_before_issue(
            db, tenant_scope=scope, document_id=draft.id, provider=provider
        )
        assert verified_revision == session.remote_revision

    issued = await ManagedDocumentService.issue(
        db,
        tenant_scope=scope,
        document_id=draft.id,
        template_storage=template_storage,
        artifact_storage=artifact_storage,
        pdf_converter=FakePdfConverter(),
        verified_remote_revision=verified_revision,
    )
    artifact = next(item for item in issued.artifacts if item.kind == "rendered_docx")
    text = _document_text(
        await artifact_storage.read(ManagedDocumentService.stored_artifact(artifact))
    )

    assert "ПРИНЯТ ТОЛЬКО ПЕРВЫЙ ЭТАП" in text
    assert "Первый этап: 3000.00" in text
    assert "Остаток: 140.00" in text
    assert "СТАРОЕ ПОЛНОЕ ПРИНЯТИЕ" not in text
    assert "{{" not in text
    if outdoor_in_first:
        assert "Первый: Установка наружного блока" in text
        assert "Второй: Установка внутреннего блока" in text
    else:
        assert "Первый: Штробление и прокладка коммуникаций" in text
        assert "Второй: Установка наружного и внутреннего блоков" in text
